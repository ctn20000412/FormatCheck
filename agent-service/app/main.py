from __future__ import annotations

import argparse
import copy
import json
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from datetime import datetime, timezone
from xml.etree import ElementTree as ET

try:
    from docx import Document
except ImportError:  # pragma: no cover - fallback for minimal Python environments.
    Document = None

try:
    from app.llm_client import call_openai_compatible_chat
    from app.llm_config import load_provider_settings
    from app.llm_config import LlmConfigStore
except ModuleNotFoundError:  # pragma: no cover - used when main.py is executed as a script.
    from llm_client import call_openai_compatible_chat
    from llm_config import load_provider_settings
    from llm_config import LlmConfigStore

BUSINESS_DIRS = {
    "checked": "\u5f85\u68c0\u6d4b\u6587\u4ef6",
    "standard": "\u89c4\u5219\u89c4\u8303\u6587\u4ef6",
    "rule": "\u89c4\u5219\u62bd\u53d6\u7ed3\u679c",
    "result": "\u68c0\u6d4b\u7ed3\u679c",
    "annotated": "\u68c0\u6d4b\u540e\u5e26\u6279\u6ce8\u5bf9\u7b56\u6e90\u6587\u4ef6",
}

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXTRACT_RULE_PROMPT_PATH = PROJECT_ROOT / "skills" / "extract-format-rules" / "references" / "extract_rule_prompt.md"
FORMAT_CHECK_PROMPT_PATH = PROJECT_ROOT / "skills" / "check-paper-format" / "references" / "format_check_prompt.md"
MAX_LLM_TEXT_CHARS = 10000
SUPPORTED_WORKFLOWS = {
    "base_format",
    "language_semantic",
    "full_check",
    "hard_format",
    "semantic_llm",
    "llm_direct",
    "hybrid",
    "compare",
}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

ET.register_namespace("w", W_NS)
ET.register_namespace("r", R_NS)
ET.register_namespace("m", M_NS)


def sanitize_folder_name(filename: str) -> str:
    stem = Path(filename).stem
    value = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]+", "_", stem)
    value = re.sub(r"_+", "_", value).strip("_")
    return value or "document"


def qn(namespace: str, tag: str) -> str:
    return f"{{{namespace}}}{tag}"


def build_request_paths(work_dir: Path, checked_filename: str, standard_filename: str) -> dict[str, Path]:
    checked_stem = Path(checked_filename).stem
    return {
        "checked_file": work_dir / BUSINESS_DIRS["checked"] / checked_filename,
        "standard_file": work_dir / BUSINESS_DIRS["standard"] / standard_filename,
        "format_rule_output": work_dir / BUSINESS_DIRS["rule"] / "format_rule.json",
        "check_result_output": work_dir / BUSINESS_DIRS["result"] / "format_check_result.json",
        "analysis_output": work_dir / BUSINESS_DIRS["result"] / "format_check_analysis.docx",
        "analysis_markdown_output": work_dir / BUSINESS_DIRS["result"] / "format_check_analysis.md",
        "llm_log_output": work_dir / BUSINESS_DIRS["result"] / "llm_interactions.jsonl",
        "annotated_output": work_dir / BUSINESS_DIRS["annotated"] / f"{checked_stem}_格式检查批注版.docx",
    }


def ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def copy_if_needed(source: Path, target: Path) -> Path:
    ensure_parent(target)
    if source.resolve() != target.resolve():
        shutil.copy2(source, target)
    return target


def write_json(path: Path, data: dict[str, Any]) -> None:
    ensure_parent(path)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def read_skill_prompt(path: Path, fallback: str) -> str:
    if path.exists():
        return path.read_text(encoding="utf-8")
    return fallback


def force_chinese_output(prompt: str) -> str:
    chinese_requirement = """

## 中文输出强制要求

必须使用中文输出。所有 JSON 字段值、分析摘要、规则名称、问题原因、规范要求、修改建议、批注内容都必须使用中文。
除 JSON schema 中固定的英文字段名外，不要输出英文说明。
如果原始规范或论文内容中有英文术语，可以保留术语本身，但解释和结论必须使用中文。
"""
    json_requirement = """

## JSON 合法性强制要求
只能输出一个完整 JSON 对象，不能输出 Markdown、解释文字或代码块。
所有字符串内部如果需要引用术语，优先使用中文引号“”，不要直接使用英文双引号。
如果必须在字符串内部使用英文双引号，必须写成 JSON 转义形式：\\\"。
输出前必须确认 JSON 可被标准 json.loads 直接解析。
"""
    return prompt.rstrip() + chinese_requirement + json_requirement


def read_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".pdf":
        return path.read_bytes()[:12000].decode("utf-8", errors="ignore")
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx_text(path: Path) -> str:
    w_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    parts: list[str] = []
    with zipfile.ZipFile(path) as package:
        xml = package.read("word/document.xml")
    root = ET.fromstring(xml)
    for paragraph in root.iter(w_ns + "p"):
        texts = [node.text or "" for node in paragraph.iter(w_ns + "t")]
        paragraph_text = "".join(texts).strip()
        if paragraph_text:
            parts.append(paragraph_text)
    return "\n".join(parts)


def length_to_pt(value: Any) -> float | None:
    if value is None:
        return None
    return round(float(value.pt), 2)


def length_to_mm(value: Any) -> float | None:
    if value is None:
        return None
    return round(float(value.cm) * 10, 2)


def element_attr(element: ET.Element | None, namespace: str, tag: str) -> str | None:
    if element is None:
        return None
    return element.get(qn(namespace, tag))


def child(element: ET.Element | None, namespace: str, tag: str) -> ET.Element | None:
    if element is None:
        return None
    return element.find(qn(namespace, tag))


def text_from_word_element(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(node.text or "" for node in element.iter(qn(W_NS, "t")))


def collect_fields(element: ET.Element | None) -> list[dict[str, Any]]:
    if element is None:
        return []
    fields: list[dict[str, Any]] = []
    for instr in element.iter(qn(W_NS, "instrText")):
        instruction = (instr.text or "").strip()
        if instruction:
            fields.append({"instruction": instruction})
    return fields


def twips_to_mm(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        return round(int(value) * 25.4 / 1440, 2)
    except ValueError:
        return None


def twips_to_pt(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        return round(int(value) / 20, 2)
    except ValueError:
        return None


def half_points_to_pt(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        return round(int(value) / 2, 2)
    except ValueError:
        return None


def word_bool(element: ET.Element | None) -> bool | None:
    if element is None:
        return None
    value = element_attr(element, W_NS, "val")
    if value is None:
        return True
    return value not in {"0", "false", "False", "off"}


def compact_dict(data: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in data.items() if value is not None and value != {} and value != []}


def merge_dicts(*items: dict[str, Any] | None) -> dict[str, Any]:
    merged: dict[str, Any] = {}
    for item in items:
        if not item:
            continue
        for key, value in item.items():
            if value is not None:
                merged[key] = value
    return merged


def read_docx_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def parse_xml_part(parts: dict[str, bytes], name: str) -> ET.Element | None:
    content = parts.get(name)
    if content is None:
        return None
    try:
        return ET.fromstring(content)
    except ET.ParseError:
        return None


def resolve_package_target(source_part: str, target: str | None) -> str | None:
    if not target:
        return None
    if target.startswith("/"):
        return target.lstrip("/")
    base_dir = posixpath.dirname(source_part)
    return posixpath.normpath(posixpath.join(base_dir, target))


def parse_relationships(parts: dict[str, bytes], rels_part: str) -> dict[str, dict[str, Any]]:
    root = parse_xml_part(parts, rels_part)
    if root is None:
        return {}
    relationships: dict[str, dict[str, Any]] = {}
    for rel in root.findall(qn(PKG_REL_NS, "Relationship")):
        rel_id = rel.get("Id")
        if not rel_id:
            continue
        relationships[rel_id] = {
            "id": rel_id,
            "type": rel.get("Type"),
            "target": rel.get("Target"),
            "target_mode": rel.get("TargetMode"),
        }
    return relationships


def parse_run_properties(r_pr: ET.Element | None) -> dict[str, Any]:
    if r_pr is None:
        return {}
    fonts = child(r_pr, W_NS, "rFonts")
    color = child(r_pr, W_NS, "color")
    underline = child(r_pr, W_NS, "u")
    return compact_dict(
        {
            "ascii_font": element_attr(fonts, W_NS, "ascii"),
            "east_asia_font": element_attr(fonts, W_NS, "eastAsia"),
            "hansi_font": element_attr(fonts, W_NS, "hAnsi"),
            "cs_font": element_attr(fonts, W_NS, "cs"),
            "font_size_pt": half_points_to_pt(element_attr(child(r_pr, W_NS, "sz"), W_NS, "val")),
            "font_size_cs_pt": half_points_to_pt(element_attr(child(r_pr, W_NS, "szCs"), W_NS, "val")),
            "bold": word_bool(child(r_pr, W_NS, "b")),
            "italic": word_bool(child(r_pr, W_NS, "i")),
            "underline": element_attr(underline, W_NS, "val") or (True if underline is not None else None),
            "color": element_attr(color, W_NS, "val"),
            "highlight": element_attr(child(r_pr, W_NS, "highlight"), W_NS, "val"),
            "vertical_align": element_attr(child(r_pr, W_NS, "vertAlign"), W_NS, "val"),
        }
    )


def parse_num_pr(num_pr: ET.Element | None, numbering_index: dict[str, dict[str, Any]] | None = None) -> dict[str, Any]:
    if num_pr is None:
        return {}
    num_id = element_attr(child(num_pr, W_NS, "numId"), W_NS, "val")
    ilvl = element_attr(child(num_pr, W_NS, "ilvl"), W_NS, "val")
    result: dict[str, Any] = compact_dict({"num_id": num_id, "ilvl": ilvl})
    if numbering_index and num_id in numbering_index:
        result.update(numbering_index[num_id])
        levels = result.get("levels") or {}
        if ilvl in levels:
            result.update(
                {
                    "num_format": levels[ilvl].get("num_format"),
                    "level_text": levels[ilvl].get("level_text"),
                }
            )
    return compact_dict(result)


def parse_paragraph_properties(
    p_pr: ET.Element | None, numbering_index: dict[str, dict[str, Any]] | None = None
) -> dict[str, Any]:
    if p_pr is None:
        return {}
    ind = child(p_pr, W_NS, "ind")
    spacing = child(p_pr, W_NS, "spacing")
    return compact_dict(
        {
            "style_id": element_attr(child(p_pr, W_NS, "pStyle"), W_NS, "val"),
            "alignment": element_attr(child(p_pr, W_NS, "jc"), W_NS, "val"),
            "outline_level": element_attr(child(p_pr, W_NS, "outlineLvl"), W_NS, "val"),
            "keep_next": word_bool(child(p_pr, W_NS, "keepNext")),
            "keep_lines": word_bool(child(p_pr, W_NS, "keepLines")),
            "page_break_before": word_bool(child(p_pr, W_NS, "pageBreakBefore")),
            "widow_control": word_bool(child(p_pr, W_NS, "widowControl")),
            "indent": compact_dict(
                {
                    "first_line_pt": twips_to_pt(element_attr(ind, W_NS, "firstLine")),
                    "hanging_pt": twips_to_pt(element_attr(ind, W_NS, "hanging")),
                    "left_pt": twips_to_pt(element_attr(ind, W_NS, "left")),
                    "right_pt": twips_to_pt(element_attr(ind, W_NS, "right")),
                }
            ),
            "spacing": compact_dict(
                {
                    "before_pt": twips_to_pt(element_attr(spacing, W_NS, "before")),
                    "after_pt": twips_to_pt(element_attr(spacing, W_NS, "after")),
                    "line": element_attr(spacing, W_NS, "line"),
                    "line_rule": element_attr(spacing, W_NS, "lineRule"),
                }
            ),
            "numbering": parse_num_pr(child(p_pr, W_NS, "numPr"), numbering_index),
        }
    )


def parse_styles(parts: dict[str, bytes]) -> dict[str, Any]:
    root = parse_xml_part(parts, "word/styles.xml")
    if root is None:
        return {"document_defaults": {}, "styles_by_id": {}}

    defaults_r_pr = child(child(child(root, W_NS, "docDefaults"), W_NS, "rPrDefault"), W_NS, "rPr")
    defaults_p_pr = child(child(child(root, W_NS, "docDefaults"), W_NS, "pPrDefault"), W_NS, "pPr")
    styles_by_id: dict[str, dict[str, Any]] = {}
    for style in root.findall(qn(W_NS, "style")):
        style_id = element_attr(style, W_NS, "styleId")
        if not style_id:
            continue
        styles_by_id[style_id] = compact_dict(
            {
                "style_id": style_id,
                "type": element_attr(style, W_NS, "type"),
                "name": element_attr(child(style, W_NS, "name"), W_NS, "val"),
                "based_on": element_attr(child(style, W_NS, "basedOn"), W_NS, "val"),
                "next": element_attr(child(style, W_NS, "next"), W_NS, "val"),
                "run_properties": parse_run_properties(child(style, W_NS, "rPr")),
                "paragraph_properties": parse_paragraph_properties(child(style, W_NS, "pPr")),
            }
        )

    return {
        "document_defaults": {
            "run_properties": parse_run_properties(defaults_r_pr),
            "paragraph_properties": parse_paragraph_properties(defaults_p_pr),
        },
        "styles_by_id": styles_by_id,
    }


def resolve_style_properties(styles: dict[str, Any], style_id: str | None) -> dict[str, dict[str, Any]]:
    defaults = styles.get("document_defaults", {})
    chain: list[dict[str, Any]] = []
    seen: set[str] = set()
    current = style_id
    styles_by_id = styles.get("styles_by_id", {})
    while current and current not in seen and current in styles_by_id:
        seen.add(current)
        style = styles_by_id[current]
        chain.insert(0, style)
        current = style.get("based_on")

    run_props = defaults.get("run_properties", {})
    paragraph_props = defaults.get("paragraph_properties", {})
    for style in chain:
        run_props = merge_dicts(run_props, style.get("run_properties"))
        paragraph_props = merge_dicts(paragraph_props, style.get("paragraph_properties"))
    return {"run_properties": run_props, "paragraph_properties": paragraph_props}


def parse_numbering(parts: dict[str, bytes]) -> dict[str, Any]:
    root = parse_xml_part(parts, "word/numbering.xml")
    if root is None:
        return {"abstract_numbers": {}, "numbers": {}, "paragraph_numbering": []}

    abstract_numbers: dict[str, dict[str, Any]] = {}
    for abstract in root.findall(qn(W_NS, "abstractNum")):
        abstract_id = element_attr(abstract, W_NS, "abstractNumId")
        if not abstract_id:
            continue
        levels: dict[str, dict[str, Any]] = {}
        for level in abstract.findall(qn(W_NS, "lvl")):
            ilvl = element_attr(level, W_NS, "ilvl")
            if ilvl is None:
                continue
            levels[ilvl] = compact_dict(
                {
                    "ilvl": ilvl,
                    "num_format": element_attr(child(level, W_NS, "numFmt"), W_NS, "val"),
                    "level_text": element_attr(child(level, W_NS, "lvlText"), W_NS, "val"),
                    "start": element_attr(child(level, W_NS, "start"), W_NS, "val"),
                    "paragraph_properties": parse_paragraph_properties(child(level, W_NS, "pPr")),
                    "run_properties": parse_run_properties(child(level, W_NS, "rPr")),
                }
            )
        abstract_numbers[abstract_id] = {"abstract_num_id": abstract_id, "levels": levels}

    numbers: dict[str, dict[str, Any]] = {}
    paragraph_numbering: list[dict[str, Any]] = []
    for number in root.findall(qn(W_NS, "num")):
        num_id = element_attr(number, W_NS, "numId")
        abstract_id = element_attr(child(number, W_NS, "abstractNumId"), W_NS, "val")
        if not num_id:
            continue
        item = compact_dict(
            {
                "num_id": num_id,
                "abstract_num_id": abstract_id,
                "levels": abstract_numbers.get(abstract_id, {}).get("levels", {}),
            }
        )
        numbers[num_id] = item
        paragraph_numbering.append(item)
    return {"abstract_numbers": abstract_numbers, "numbers": numbers, "paragraph_numbering": paragraph_numbering}


def parse_sections(document_root: ET.Element | None, relationships: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    if document_root is None:
        return []
    sections: list[dict[str, Any]] = []
    for index, sect_pr in enumerate(document_root.iter(qn(W_NS, "sectPr")), start=1):
        page_size = child(sect_pr, W_NS, "pgSz")
        page_margin = child(sect_pr, W_NS, "pgMar")
        header_refs = []
        for ref in sect_pr.findall(qn(W_NS, "headerReference")):
            rel_id = element_attr(ref, R_NS, "id")
            header_refs.append(
                compact_dict(
                    {
                        "type": element_attr(ref, W_NS, "type"),
                        "relationship_id": rel_id,
                        "target": relationships.get(rel_id, {}).get("target") if rel_id else None,
                    }
                )
            )
        footer_refs = []
        for ref in sect_pr.findall(qn(W_NS, "footerReference")):
            rel_id = element_attr(ref, R_NS, "id")
            footer_refs.append(
                compact_dict(
                    {
                        "type": element_attr(ref, W_NS, "type"),
                        "relationship_id": rel_id,
                        "target": relationships.get(rel_id, {}).get("target") if rel_id else None,
                    }
                )
            )
        sections.append(
            compact_dict(
                {
                    "index": index,
                    "page_width_mm": twips_to_mm(element_attr(page_size, W_NS, "w")),
                    "page_height_mm": twips_to_mm(element_attr(page_size, W_NS, "h")),
                    "orientation": element_attr(page_size, W_NS, "orient"),
                    "top_margin_mm": twips_to_mm(element_attr(page_margin, W_NS, "top")),
                    "bottom_margin_mm": twips_to_mm(element_attr(page_margin, W_NS, "bottom")),
                    "left_margin_mm": twips_to_mm(element_attr(page_margin, W_NS, "left")),
                    "right_margin_mm": twips_to_mm(element_attr(page_margin, W_NS, "right")),
                    "header_distance_mm": twips_to_mm(element_attr(page_margin, W_NS, "header")),
                    "footer_distance_mm": twips_to_mm(element_attr(page_margin, W_NS, "footer")),
                    "gutter_mm": twips_to_mm(element_attr(page_margin, W_NS, "gutter")),
                    "columns": element_attr(child(sect_pr, W_NS, "cols"), W_NS, "num"),
                    "header_references": header_refs,
                    "footer_references": footer_refs,
                }
            )
        )
    return sections


def parse_paragraphs(
    document_root: ET.Element | None,
    styles: dict[str, Any],
    numbering: dict[str, Any],
) -> list[dict[str, Any]]:
    if document_root is None:
        return []
    numbering_index = numbering.get("numbers", {})
    paragraphs: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document_root.iter(qn(W_NS, "p")), start=1):
        p_pr = child(paragraph, W_NS, "pPr")
        direct_p_props = parse_paragraph_properties(p_pr, numbering_index)
        style_id = direct_p_props.get("style_id")
        style_props = resolve_style_properties(styles, style_id)
        numbering_props = direct_p_props.get("numbering", {})
        runs: list[dict[str, Any]] = []
        for run_index, run in enumerate(paragraph.findall(qn(W_NS, "r")), start=1):
            run_text = text_from_word_element(run)
            direct_run_props = parse_run_properties(child(run, W_NS, "rPr"))
            drawing_rel_ids = [
                element_attr(blip, R_NS, "embed")
                for blip in run.iter(qn(A_NS, "blip"))
                if element_attr(blip, R_NS, "embed")
            ]
            fields = collect_fields(run)
            runs.append(
                compact_dict(
                    {
                        "index": run_index,
                        "text": run_text,
                        "direct_properties": direct_run_props,
                        "effective_properties": merge_dicts(style_props.get("run_properties"), direct_run_props),
                        "drawing_relationship_ids": drawing_rel_ids,
                        "fields": fields,
                    }
                )
            )

        paragraph_item = compact_dict(
            {
                "index": index,
                "text": text_from_word_element(paragraph),
                "style_id": style_id,
                "style_name": styles.get("styles_by_id", {}).get(style_id, {}).get("name"),
                "direct_paragraph_properties": direct_p_props,
                "effective_paragraph_properties": merge_dicts(
                    style_props.get("paragraph_properties"), direct_p_props
                ),
                "numbering": numbering_props,
                "runs": runs,
            }
        )
        paragraph_item["content_roles"] = classify_paragraph_content_roles(paragraph_item)
        paragraphs.append(compact_dict(paragraph_item))
    return paragraphs


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def infer_heading_role(paragraph: dict[str, Any]) -> str | None:
    style_id = compact_text(paragraph.get("style_id"))
    style_name = compact_text(paragraph.get("style_name"))
    text = str(paragraph.get("text") or "").strip()
    effective = paragraph.get("effective_paragraph_properties") or {}
    outline_level = str(effective.get("outline_level") or "")
    numbering = paragraph.get("numbering") or {}
    ilvl = str(numbering.get("ilvl") or "")

    heading_sources = [style_id, style_name]
    for level in range(1, 10):
        if any(f"heading{level}" in source or f"标题{level}" in source for source in heading_sources):
            return f"heading_{level}"
    if outline_level.isdigit():
        return f"heading_{int(outline_level) + 1}"
    if ilvl.isdigit() and re.match(r"^(第.+章|\d+(?:\.\d+){0,2})(\s+|[、.．])", text):
        return f"heading_{int(ilvl) + 1}"
    if re.match(r"^第[一二三四五六七八九十百千万\d]+章\b", text):
        return "heading_1"
    if re.match(r"^\d+\.\d+\.\d+(\s+|[、.．])", text):
        return "heading_3"
    if re.match(r"^\d+\.\d+(\s+|[、.．])", text):
        return "heading_2"
    if re.match(r"^\d+(\s+|[、.．])", text) and len(text) <= 80:
        return "heading_1"
    return None


def classify_paragraph_content_roles(paragraph: dict[str, Any]) -> list[str]:
    text = str(paragraph.get("text") or "").strip()
    normalized = compact_text(text)
    roles: list[str] = []
    if not text:
        return ["blank"]

    if normalized in {"目录", "目次"} or re.match(r"^.{2,80}[.\s·…]{2,}\d+$", text):
        roles.append("toc")
    if normalized in {"摘要", "中文摘要", "abstract"} or "关键词" in text or "key words" in normalized:
        roles.append("abstract")
    if normalized.startswith("参考文献") or re.match(r"^\[\d+\]", text):
        roles.append("reference")
    if re.match(r"^(表|table)\s*\d+", text, flags=re.IGNORECASE):
        roles.append("table_caption")
    if re.match(r"^(图|figure|fig\.)\s*\d+", text, flags=re.IGNORECASE):
        roles.append("figure_caption")
    if re.search(r"[\(\uff08]\s*\d+(?:[.-]\d+)?\s*[\)\uff09]\s*$", text) or any(
        field.get("instruction", "").strip().startswith("EQ")
        for run in paragraph.get("runs") or []
        for field in run.get("fields") or []
        if isinstance(field, dict)
    ):
        roles.append("formula")

    heading_role = infer_heading_role(paragraph)
    if heading_role:
        roles.extend(["heading", heading_role])

    if not roles:
        roles.append("body")
    return sorted(set(roles))


def parse_table_borders(tbl_pr: ET.Element | None) -> dict[str, Any]:
    borders = child(tbl_pr, W_NS, "tblBorders")
    if borders is None:
        return {}
    result: dict[str, Any] = {}
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = child(borders, W_NS, name)
        if border is not None:
            result[name] = compact_dict(
                {
                    "val": element_attr(border, W_NS, "val"),
                    "size": element_attr(border, W_NS, "sz"),
                    "color": element_attr(border, W_NS, "color"),
                }
            )
    return result


def parse_tables(document_root: ET.Element | None) -> list[dict[str, Any]]:
    if document_root is None:
        return []
    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(document_root.iter(qn(W_NS, "tbl")), start=1):
        rows = []
        max_columns = 0
        for row_index, row in enumerate(table.findall(qn(W_NS, "tr")), start=1):
            cells = []
            for cell_index, cell in enumerate(row.findall(qn(W_NS, "tc")), start=1):
                cells.append({"index": cell_index, "text": text_from_word_element(cell)})
            max_columns = max(max_columns, len(cells))
            rows.append({"index": row_index, "cells": cells})
        tbl_pr = child(table, W_NS, "tblPr")
        tbl_width = child(tbl_pr, W_NS, "tblW")
        tables.append(
            compact_dict(
                {
                    "index": table_index,
                    "row_count": len(rows),
                    "column_count": max_columns,
                    "width": compact_dict(
                        {
                            "type": element_attr(tbl_width, W_NS, "type"),
                            "value": element_attr(tbl_width, W_NS, "w"),
                        }
                    ),
                    "alignment": element_attr(child(tbl_pr, W_NS, "jc"), W_NS, "val"),
                    "borders": parse_table_borders(tbl_pr),
                    "rows": rows,
                    "preview_rows": [[cell["text"] for cell in row["cells"]] for row in rows[:5]],
                }
            )
        )
    return tables


def parse_related_text_parts(
    parts: dict[str, bytes],
    document_relationships: dict[str, dict[str, Any]],
    rel_type_suffix: str,
    source_part: str,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for rel_id, rel in document_relationships.items():
        rel_type = rel.get("type") or ""
        if not rel_type.endswith(rel_type_suffix):
            continue
        part_name = resolve_package_target(source_part, rel.get("target"))
        root = parse_xml_part(parts, part_name or "")
        if root is None:
            continue
        items.append(
            compact_dict(
                {
                    "relationship_id": rel_id,
                    "target": rel.get("target"),
                    "part_name": part_name,
                    "text": text_from_word_element(root),
                    "fields": collect_fields(root),
                }
            )
        )
    return items


def parse_note_part(parts: dict[str, bytes], part_name: str, item_tag: str) -> list[dict[str, Any]]:
    root = parse_xml_part(parts, part_name)
    if root is None:
        return []
    notes: list[dict[str, Any]] = []
    for note in root.findall(qn(W_NS, item_tag)):
        note_id = element_attr(note, W_NS, "id")
        notes.append(compact_dict({"id": note_id, "text": text_from_word_element(note)}))
    return notes


def parse_images(
    parts: dict[str, bytes],
    document_root: ET.Element | None,
    document_relationships: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    images: list[dict[str, Any]] = []
    if document_root is None:
        return images
    seen: set[tuple[str, str]] = set()
    for blip in document_root.iter(qn(A_NS, "blip")):
        rel_id = element_attr(blip, R_NS, "embed") or element_attr(blip, R_NS, "link")
        if not rel_id:
            continue
        rel = document_relationships.get(rel_id, {})
        target = rel.get("target")
        key = (rel_id, target or "")
        if key in seen:
            continue
        seen.add(key)
        part_name = resolve_package_target("word/document.xml", target)
        images.append(
            compact_dict(
                {
                    "relationship_id": rel_id,
                    "target": target,
                    "part_name": part_name,
                    "content_type": "image" if part_name in parts else None,
                }
            )
        )
    return images


def parse_hyperlinks(
    document_root: ET.Element | None,
    document_relationships: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    if document_root is None:
        return []
    hyperlinks: list[dict[str, Any]] = []
    for index, hyperlink in enumerate(document_root.iter(qn(W_NS, "hyperlink")), start=1):
        rel_id = element_attr(hyperlink, R_NS, "id")
        hyperlinks.append(
            compact_dict(
                {
                    "index": index,
                    "relationship_id": rel_id,
                    "anchor": element_attr(hyperlink, W_NS, "anchor"),
                    "target": document_relationships.get(rel_id, {}).get("target") if rel_id else None,
                    "text": text_from_word_element(hyperlink),
                }
            )
        )
    return hyperlinks


def _legacy_build_docx_analysis_evidence(path: Path) -> dict[str, Any]:
    if path.suffix.lower() != ".docx":
        return {
            "parser": "plain-text-reader",
            "file_name": path.name,
            "note": "non-docx file; only plain text evidence is available.",
        }
    if Document is None:
        return {
            "parser": "zip-xml-fallback",
            "file_name": path.name,
            "note": "python-docx is unavailable; using zip XML fallback evidence.",
            "full_text": read_docx_text(path),
        }

    document = Document(str(path))
    paragraphs: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_format = paragraph.paragraph_format
        runs = []
        for run in paragraph.runs:
            run_text = run.text.strip()
            if not run_text:
                continue
            runs.append(
                {
                    "text": run_text[:80],
                    "font_name": run.font.name,
                    "font_size_pt": length_to_pt(run.font.size),
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": bool(run.underline),
                }
            )
            if len(runs) >= 8:
                break
        paragraphs.append(
            {
                "index": index,
                "text": text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                "first_line_indent_pt": length_to_pt(paragraph_format.first_line_indent),
                "left_indent_pt": length_to_pt(paragraph_format.left_indent),
                "right_indent_pt": length_to_pt(paragraph_format.right_indent),
                "line_spacing": paragraph_format.line_spacing,
                "space_before_pt": length_to_pt(paragraph_format.space_before),
                "space_after_pt": length_to_pt(paragraph_format.space_after),
                "runs": runs,
            }
        )

    sections = []
    for index, section in enumerate(document.sections, start=1):
        sections.append(
            {
                "index": index,
                "page_width_mm": length_to_mm(section.page_width),
                "page_height_mm": length_to_mm(section.page_height),
                "top_margin_mm": length_to_mm(section.top_margin),
                "bottom_margin_mm": length_to_mm(section.bottom_margin),
                "left_margin_mm": length_to_mm(section.left_margin),
                "right_margin_mm": length_to_mm(section.right_margin),
                "header_distance_mm": length_to_mm(section.header_distance),
                "footer_distance_mm": length_to_mm(section.footer_distance),
            }
        )

    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        preview_rows = []
        for row in table.rows[:5]:
            preview_rows.append([cell.text.strip() for cell in row.cells])
        tables.append(
            {
                "index": table_index,
                "row_count": len(table.rows),
                "column_count": len(table.columns),
                "preview_rows": preview_rows,
            }
        )

    return {
        "parser": "python-docx",
        "file_name": path.name,
        "paragraph_count": len(paragraphs),
        "section_count": len(sections),
        "table_count": len(tables),
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
    }


def build_docx_analysis_evidence(path: Path) -> dict[str, Any]:
    if path.suffix.lower() != ".docx":
        return {
            "parser": "plain-text-reader",
            "file_name": path.name,
            "note": "Only .docx files expose WordprocessingML formatting evidence.",
        }

    parts = read_docx_parts(path)
    document_root = parse_xml_part(parts, "word/document.xml")
    document_relationships = parse_relationships(parts, "word/_rels/document.xml.rels")
    styles = parse_styles(parts)
    numbering = parse_numbering(parts)
    paragraphs = parse_paragraphs(document_root, styles, numbering)
    sections = parse_sections(document_root, document_relationships)
    tables = parse_tables(document_root)
    headers = parse_related_text_parts(parts, document_relationships, "/header", "word/document.xml")
    footers = parse_related_text_parts(parts, document_relationships, "/footer", "word/document.xml")
    fields = collect_fields(document_root)
    footnotes = parse_note_part(parts, "word/footnotes.xml", "footnote")
    endnotes = parse_note_part(parts, "word/endnotes.xml", "endnote")
    images = parse_images(parts, document_root, document_relationships)
    hyperlinks = parse_hyperlinks(document_root, document_relationships)

    python_docx_summary: dict[str, Any] = {}
    if Document is not None:
        try:
            document = Document(str(path))
            python_docx_summary = {
                "paragraph_count": len(document.paragraphs),
                "section_count": len(document.sections),
                "table_count": len(document.tables),
            }
        except Exception as exc:  # pragma: no cover - depends on malformed source files.
            python_docx_summary = {"error": str(exc)}

    return {
        "parser": "python-docx+xml",
        "file_name": path.name,
        "package_parts_present": sorted(parts.keys()),
        "coverage": {
            "source": "docx zip package WordprocessingML",
            "includes": [
                "document body paragraphs and runs",
                "direct and style-derived run properties",
                "direct and style-derived paragraph properties",
                "numbering definitions",
                "section page setup and margins",
                "tables and table borders",
                "headers and footers",
                "fields",
                "footnotes and endnotes",
                "images",
                "hyperlinks",
            ],
            "limits": [
                "does not calculate rendered page numbers without a Word layout engine",
                "does not OCR embedded images",
            ],
        },
        "python_docx_summary": python_docx_summary,
        "full_text": "\n".join(str(paragraph.get("text") or "") for paragraph in paragraphs),
        "paragraph_count": len(paragraphs),
        "section_count": len(sections),
        "table_count": len(tables),
        "styles": styles,
        "numbering": numbering,
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
        "headers_footers": {"headers": headers, "footers": footers},
        "fields": fields,
        "footnotes": footnotes,
        "endnotes": endnotes,
        "images": images,
        "hyperlinks": hyperlinks,
    }


def extract_json_text(content: str) -> str:
    value = content.strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*", "", value)
        value = re.sub(r"\s*```$", "", value)
    start = value.find("{")
    end = value.rfind("}")
    if start >= 0 and end > start:
        value = value[start : end + 1]
    return value


def remove_trailing_json_commas(value: str) -> str:
    return re.sub(r",\s*([}\]])", r"\1", value)


def escape_unescaped_quotes_in_json_strings(value: str) -> str:
    repaired: list[str] = []
    in_string = False
    escaped = False
    for index, char in enumerate(value):
        if not in_string:
            if char == '"':
                in_string = True
            repaired.append(char)
            continue

        if escaped:
            repaired.append(char)
            escaped = False
            continue

        if char == "\\":
            repaired.append(char)
            escaped = True
            continue

        if char == '"':
            next_index = index + 1
            while next_index < len(value) and value[next_index].isspace():
                next_index += 1
            next_char = value[next_index] if next_index < len(value) else ""
            if next_char in {":", ",", "}", "]", ""}:
                in_string = False
                repaired.append(char)
            else:
                repaired.append('\\"')
            continue

        repaired.append(char)
    return "".join(repaired)


def format_json_error_excerpt(value: str, error: json.JSONDecodeError) -> str:
    lines = value.splitlines()
    start = max(0, error.lineno - 3)
    end = min(len(lines), error.lineno + 2)
    excerpt = []
    for line_number in range(start, end):
        excerpt.append(f"{line_number + 1}: {lines[line_number]}")
    return "\n".join(excerpt)


def parse_llm_json(content: str, context: str = "LLM") -> dict[str, Any]:
    value = extract_json_text(content)
    candidates = [
        value,
        remove_trailing_json_commas(value),
        escape_unescaped_quotes_in_json_strings(remove_trailing_json_commas(value)),
    ]
    last_error: json.JSONDecodeError | None = None
    for candidate in candidates:
        try:
            return json.loads(candidate)
        except json.JSONDecodeError as exc:
            last_error = exc

    assert last_error is not None
    excerpt = format_json_error_excerpt(value, last_error)
    raise ValueError(
        f"{context} failed to parse LLM JSON at line {last_error.lineno}, "
        f"column {last_error.colno}: {last_error.msg}\nExcerpt:\n{excerpt}"
    ) from last_error


def split_text_chunks(text: str, max_chars: int = MAX_LLM_TEXT_CHARS) -> list[str]:
    if len(text) <= max_chars:
        return [text]

    chunks: list[str] = []
    current: list[str] = []
    current_length = 0
    for paragraph in text.splitlines():
        paragraph_with_newline = paragraph + "\n"
        if current and current_length + len(paragraph_with_newline) > max_chars:
            chunks.append("".join(current).rstrip())
            current = []
            current_length = 0

        if len(paragraph_with_newline) > max_chars:
            if current:
                chunks.append("".join(current).rstrip())
                current = []
                current_length = 0
            for start in range(0, len(paragraph_with_newline), max_chars):
                chunks.append(paragraph_with_newline[start : start + max_chars].rstrip())
            continue

        current.append(paragraph_with_newline)
        current_length += len(paragraph_with_newline)

    if current:
        chunks.append("".join(current).rstrip())
    return [chunk for chunk in chunks if chunk]


def aggregate_check_payloads(
    payloads: list[dict[str, Any]],
    checked_rule_count: int,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    for chunk_index, payload in enumerate(payloads, start=1):
        for issue in payload.get("issues", []):
            if isinstance(issue, dict):
                normalized_issue = dict(issue)
                normalized_issue.setdefault("chunk_index", chunk_index)
                issues.append(normalized_issue)

    by_category: dict[str, int] = {}
    by_severity: dict[str, int] = {}
    need_manual_confirmation = 0
    for issue in issues:
        category = str(issue.get("category") or "unknown")
        severity = str(issue.get("severity") or "unknown")
        by_category[category] = by_category.get(category, 0) + 1
        by_severity[severity] = by_severity.get(severity, 0) + 1
        if issue.get("need_manual_confirmation") is True:
            need_manual_confirmation += 1

    return (
        {
            "total_issues": len(issues),
            "by_category": by_category,
            "by_severity": by_severity,
            "need_manual_confirmation": need_manual_confirmation,
            "checked_rule_count": checked_rule_count,
        },
        issues,
    )


def build_document_overview(text: str, chunks: list[str]) -> dict[str, Any]:
    return {
        "total_chars": len(text),
        "chunk_count": len(chunks),
        "chunk_char_lengths": [len(chunk) for chunk in chunks],
        "document_start": text[:300],
        "document_end": text[-500:],
    }


def build_location_display(issue: dict[str, Any]) -> str:
    location = issue.get("location")
    if isinstance(location, str):
        return location
    if not isinstance(location, dict):
        return str(issue.get("position") or issue.get("anchor_text") or "")
    parts = []
    for key in ("page", "section", "paragraph_index", "object_type"):
        value = location.get(key)
        if value not in (None, ""):
            label = {
                "page": "\u9875\u7801",
                "section": "\u7ae0\u8282",
                "paragraph_index": "\u6bb5\u843d",
                "object_type": "\u5bf9\u8c61",
            }[key]
            parts.append(f"{label}: {value}")
    quote = location.get("quote") or issue.get("anchor_text")
    if quote:
        parts.append(f"\u951a\u70b9: {quote}")
    return "；".join(parts)


def normalize_issue(issue: dict[str, Any], index: int) -> dict[str, Any]:
    normalized = dict(issue)
    normalized.setdefault("issue_id", normalized.get("id") or f"I{index:03d}")
    normalized.setdefault("id", normalized["issue_id"])
    location = normalized.get("location")
    if not isinstance(location, dict):
        normalized["location"] = {
            "page": None,
            "section": str(location or ""),
            "paragraph_index": None,
            "object_type": "other",
            "quote": normalized.get("anchor_text") or "",
        }
    normalized.setdefault("anchor_text", normalized.get("location", {}).get("quote", ""))
    normalized.setdefault("comment_text", build_comment_text(normalized))
    normalized["location_display"] = build_location_display(normalized)
    return normalized


def build_comment_text(issue: dict[str, Any]) -> str:
    if issue.get("comment_text"):
        return str(issue["comment_text"])
    reason = issue.get("reason") or issue.get("error_reason") or issue.get("problem") or "未说明"
    expected = issue.get("expected") or issue.get("requirement") or issue.get("rule") or "未说明"
    suggestion = issue.get("suggestion") or issue.get("fix_suggestion") or "璇锋寜瑙勮寖淇敼"
    return f"错误原因：{reason}\n规范要求：{expected}\n修改建议：{suggestion}"


def normalize_workflow(value: Any) -> str:
    workflow = str(value or "full_check").strip() or "full_check"
    aliases = {
        "hard_format": "base_format",
        "semantic_llm": "language_semantic",
        "llm_direct": "language_semantic",
        "hybrid": "full_check",
        "compare": "full_check",
    }
    workflow = aliases.get(workflow, workflow)
    if workflow not in SUPPORTED_WORKFLOWS:
        raise ValueError(f"Unsupported workflow: {workflow}")
    return workflow


def recompute_statistics(issues: list[dict[str, Any]], checked_rule_count: int) -> dict[str, Any]:
    by_category: dict[str, int] = {}
    by_severity: dict[str, int] = {}
    need_manual_confirmation = 0
    for issue in issues:
        category = str(issue.get("category") or issue.get("rule_category") or "unknown")
        severity = str(issue.get("severity") or "medium")
        by_category[category] = by_category.get(category, 0) + 1
        by_severity[severity] = by_severity.get(severity, 0) + 1
        if issue.get("need_manual_confirmation"):
            need_manual_confirmation += 1
    return {
        "total_issues": len(issues),
        "by_category": by_category,
        "by_severity": by_severity,
        "need_manual_confirmation": need_manual_confirmation,
        "checked_rule_count": checked_rule_count,
    }


def set_result_workflow(
    format_rule: dict[str, Any],
    check_result: dict[str, Any],
    statistics: dict[str, Any],
    workflow: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule["workflow"] = workflow
    check_result["workflow"] = workflow
    check_result.setdefault("paper_analysis", {})["workflow"] = workflow
    statistics["workflow"] = workflow
    check_result["statistics"] = statistics
    return format_rule, check_result, statistics


def infer_check_method(rule: dict[str, Any]) -> str:
    text = json.dumps(rule, ensure_ascii=False).lower()
    deterministic_keywords = [
        "font",
        "size",
        "margin",
        "spacing",
        "indent",
        "alignment",
        "page",
        "table",
        "heading",
        "\u5b57\u4f53",
        "\u5b57\u53f7",
        "\u9875\u8fb9\u8ddd",
        "\u884c\u8ddd",
        "\u7f29\u8fdb",
        "\u5bf9\u9f50",
        "\u9875\u7801",
        "\u8868\u683c",
        "\u6807\u9898",
        "\u516c\u5f0f",
        "formula",
    ]
    return "python" if any(keyword in text for keyword in deterministic_keywords) else "llm"


def normalize_check_method(value: Any) -> str:
    method = str(value or "").strip().lower()
    if method in {"python", "deterministic", "hard_format", "rule_engine"}:
        return "python"
    if method in {"llm", "llm_semantic", "semantic", "semantic_llm"}:
        return "llm"
    return method or "llm"


def normalize_format_rules(format_rule: dict[str, Any]) -> list[dict[str, Any]]:
    normalized_rules = []
    for index, rule in enumerate(format_rule.get("rules", []), start=1):
        if not isinstance(rule, dict):
            continue
        normalized = dict(rule)
        normalized.setdefault("rule_id", normalized.get("id") or f"R{index:03d}")
        normalized.setdefault("id", normalized["rule_id"])
        normalized["check_method"] = normalize_check_method(normalized.get("check_method") or infer_check_method(normalized))
        normalized_rules.append(normalized)
    return normalized_rules


def filter_rules_by_check_method(rules: list[dict[str, Any]], checker: str) -> list[dict[str, Any]]:
    expected = normalize_check_method(checker)
    return [
        rule
        for rule in normalize_format_rules({"rules": rules})
        if normalize_check_method(rule.get("check_method")) == expected
    ]


def build_deterministic_issue(rule: dict[str, Any], paragraph: dict[str, Any], actual: str, expected: str) -> dict[str, Any]:
    paragraph_index = paragraph.get("index")
    quote = str(paragraph.get("text") or "")[:60]
    rule_id = rule.get("rule_id") or rule.get("id") or "R000"
    return normalize_issue(
        {
            "issue_id": f"D-{rule_id}-{paragraph_index}",
            "rule_id": rule_id,
            "category": rule.get("category") or "format",
            "severity": rule.get("severity") or "medium",
            "reason": f"\u7a0b\u5e8f\u8bfb\u53d6\u5230\u7684\u5b9e\u9645\u683c\u5f0f\u4e3a {actual}",
            "expected": f"\u89c4\u8303\u8981\u6c42\u4e3a {expected}",
            "suggestion": f"\u8bf7\u5c06\u8be5\u4f4d\u7f6e\u8c03\u6574\u4e3a {expected}",
            "need_manual_confirmation": False,
            "source": "deterministic_checker",
            "location": {
                "page": None,
                "section": paragraph.get("section_index"),
                "paragraph_index": paragraph_index,
                "object_type": "paragraph",
                "quote": quote,
            },
            "anchor_text": quote,
        },
        1,
    )


def build_python_issue(
    rule: dict[str, Any],
    category: str,
    reason: str,
    expected: str,
    suggestion: str,
    location: dict[str, Any] | None = None,
    anchor_text: str | None = None,
    manual: bool = False,
) -> dict[str, Any]:
    rule_id = rule.get("rule_id") or rule.get("id") or "R000"
    location = location or {}
    location.setdefault("page", None)
    location.setdefault("section", None)
    location.setdefault("paragraph_index", None)
    location.setdefault("object_type", category)
    quote = anchor_text or str(location.get("quote") or "")[:60]
    location.setdefault("quote", quote)
    return {
        "issue_id": f"D-{rule_id}-{category}-{location.get('paragraph_index') or location.get('section') or location.get('object_type')}",
        "rule_id": rule_id,
        "category": category,
        "severity": rule.get("severity") or "medium",
        "reason": reason,
        "expected": expected,
        "suggestion": suggestion,
        "need_manual_confirmation": manual,
        "source": "python_hard_checker",
        "location": location,
        "anchor_text": quote,
    }


def contains_cjk(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in text)


def contains_latin(text: str) -> bool:
    return any(("A" <= char <= "Z") or ("a" <= char <= "z") for char in text)


def expected_font_script(rule_text: str, expected_font: str) -> str:
    english_fonts = {"Times New Roman", "Arial", "Calibri", "Cambria"}
    chinese_fonts = {"\u5b8b\u4f53", "SimSun", "\u9ed1\u4f53", "\u6977\u4f53", "KaiTi", "FangSong"}
    if expected_font in english_fonts or "english" in rule_text.lower() or "\u82f1\u6587" in rule_text:
        return "latin"
    if expected_font in chinese_fonts or "\u4e2d\u6587" in rule_text or "\u6c49\u5b57" in rule_text:
        return "cjk"
    return "any"


def select_run_font_for_rule(run: dict[str, Any], rule_script: str) -> tuple[str | None, bool]:
    text = str(run.get("text") or "")
    effective = run.get("effective_properties") or {}
    if rule_script == "latin":
        return effective.get("ascii_font") or effective.get("hansi_font"), contains_latin(text)
    if rule_script == "cjk":
        return effective.get("east_asia_font"), contains_cjk(text)
    if contains_latin(text):
        return effective.get("ascii_font") or effective.get("hansi_font"), True
    if contains_cjk(text):
        return effective.get("east_asia_font"), True
    return effective.get("east_asia_font") or effective.get("ascii_font") or effective.get("hansi_font"), bool(text)


def rule_scope_text(rule: dict[str, Any]) -> str:
    return " ".join(
        str(rule.get(key) or "")
        for key in ("scope", "rule_name", "category", "description", "requirement", "check_logic")
    )


def normalize_rule_target_roles(rule: dict[str, Any]) -> set[str]:
    text = compact_text(rule_scope_text(rule))
    roles: set[str] = set()
    if not text:
        return roles
    if any(keyword in text for keyword in ("整篇", "全文", "whole", "entiredocument", "alldocument")):
        roles.add("all")
    if any(keyword in text for keyword in ("正文", "body")):
        roles.add("body")
    if any(keyword in text for keyword in ("中文摘要", "英文摘要", "摘要", "abstract", "关键词", "keywords")):
        roles.add("abstract")
    if any(keyword in text for keyword in ("目录", "toc", "tableofcontents")):
        roles.add("toc")
    if any(keyword in text for keyword in ("参考文献", "bibliography", "references")):
        roles.add("reference")
    if any(keyword in text for keyword in ("一级标题", "章标题", "heading1", "first-levelheading", "level1heading")):
        roles.add("heading_1")
    if any(keyword in text for keyword in ("二级标题", "节标题", "heading2", "second-levelheading", "level2heading")):
        roles.add("heading_2")
    if any(keyword in text for keyword in ("三级标题", "条标题", "heading3", "third-levelheading", "level3heading")):
        roles.add("heading_3")
    if "标题" in text or "heading" in text:
        roles.add("heading")
    if any(keyword in text for keyword in ("表格标题", "表题", "tablecaption")):
        roles.add("table_caption")
    elif any(keyword in text for keyword in ("表格", "表内", "table")):
        roles.add("table")
    if any(keyword in text for keyword in ("图片标题", "图题", "figurecaption", "figcaption")):
        roles.add("figure_caption")
    elif any(keyword in text for keyword in ("图片", "插图", "figure", "image")):
        roles.add("figure")
    if any(keyword in text for keyword in ("公式", "formula", "equation")):
        roles.add("formula")
    if any(keyword in text for keyword in ("页眉", "header")):
        roles.add("header")
    if any(keyword in text for keyword in ("页脚", "footer", "页码", "pagenumber")):
        roles.add("footer")
    return roles


def paragraph_matches_rule_scope(paragraph: dict[str, Any], rule_roles: set[str]) -> bool:
    if not rule_roles or "all" in rule_roles:
        return True
    paragraph_roles = set(paragraph.get("content_roles") or classify_paragraph_content_roles(paragraph))
    if "heading" in rule_roles and "heading" in paragraph_roles:
        return True
    if rule_roles.intersection(paragraph_roles):
        return True
    # Table and figure object rules are handled by dedicated checkers. Do not apply them to every paragraph.
    non_paragraph_roles = {"table", "figure", "header", "footer"}
    if rule_roles and rule_roles.issubset(non_paragraph_roles):
        return False
    return False


def paragraphs_matching_rule_scope(rule: dict[str, Any], paragraphs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rule_roles = normalize_rule_target_roles(rule)
    return [paragraph for paragraph in paragraphs if paragraph_matches_rule_scope(paragraph, rule_roles)]


FONT_SIZE_NAME_TO_PT = {
    "初号": 42.0,
    "小初": 36.0,
    "一号": 26.0,
    "小一": 24.0,
    "二号": 22.0,
    "小二": 18.0,
    "三号": 16.0,
    "小三": 15.0,
    "四号": 14.0,
    "小四": 12.0,
    "五号": 10.5,
    "小五": 9.0,
    "六号": 7.5,
    "小六": 6.5,
    "七号": 5.5,
    "八号": 5.0,
}


def parse_expected_font_size_pt(rule_text: str) -> float | None:
    for name in sorted(FONT_SIZE_NAME_TO_PT, key=len, reverse=True):
        if name in rule_text:
            return FONT_SIZE_NAME_TO_PT[name]
    match = re.search(r"(\d+(?:\.\d+)?)\s*(?:pt|磅)", rule_text, flags=re.IGNORECASE)
    return float(match.group(1)) if match else None


def parse_expected_alignment(rule_text: str) -> str | None:
    normalized = rule_text.lower()
    if "两端对齐" in rule_text or "justify" in normalized:
        return "both"
    if "居中" in rule_text or "center" in normalized or "centre" in normalized:
        return "center"
    if "右对齐" in rule_text or "居右" in rule_text or "right" in normalized:
        return "right"
    if "左对齐" in rule_text or "居左" in rule_text or "left" in normalized:
        return "left"
    return None


def same_number(actual: Any, expected: float, tolerance: float = 0.15) -> bool:
    try:
        return abs(float(actual) - expected) <= tolerance
    except (TypeError, ValueError):
        return False


def format_pt(value: Any) -> str:
    try:
        return f"{float(value):g} pt"
    except (TypeError, ValueError):
        return str(value)


def parse_expected_mm_values(rule_text: str) -> dict[str, float]:
    values: dict[str, float] = {}
    for key, aliases in {
        "top_margin_mm": ["上", "上边距", "top"],
        "bottom_margin_mm": ["下", "下边距", "bottom"],
        "left_margin_mm": ["左", "左边距", "left"],
        "right_margin_mm": ["右", "右边距", "right"],
        "header_distance_mm": ["页眉", "header"],
        "footer_distance_mm": ["页脚", "footer"],
    }.items():
        for alias in aliases:
            match = re.search(rf"{re.escape(alias)}\s*(?:边距|距)?\s*(?:为|[:：])?\s*(\d+(?:\.\d+)?)\s*mm", rule_text, re.IGNORECASE)
            if match:
                values[key] = float(match.group(1))
                break
    return values


def parse_expected_line_value(rule_text: str) -> int | None:
    if "1.5" in rule_text or "1．5" in rule_text:
        return 360
    if "单倍" in rule_text or "single" in rule_text.lower():
        return 240
    return None


def parse_expected_spacing_points(rule_text: str) -> dict[str, float]:
    values: dict[str, float] = {}
    if "段前" in rule_text:
        match = re.search(r"段前\s*(?:、?段后)?\s*(?:各)?\s*(?:设为|为)?\s*(\d+(?:\.\d+)?)\s*(行|pt|磅)?", rule_text)
        if match:
            value = float(match.group(1))
            values["before_pt"] = value * 12 if match.group(2) == "行" else value
    if "段后" in rule_text:
        match = re.search(r"段后\s*(?:各)?\s*(?:设为|为)?\s*(\d+(?:\.\d+)?)\s*(行|pt|磅)?", rule_text)
        if match:
            value = float(match.group(1))
            values["after_pt"] = value * 12 if match.group(2) == "行" else value
    if "段前、段后无空行" in rule_text or "段前段后间距是否为0" in rule_text or "段前0" in rule_text:
        values.setdefault("before_pt", 0.0)
    if "段前、段后无空行" in rule_text or "段前段后间距是否为0" in rule_text or "段后0" in rule_text:
        values.setdefault("after_pt", 0.0)
    return values


def parse_expected_first_line_indent_pt(rule_text: str) -> float | None:
    match = re.search(r"首行缩进\s*(\d+(?:\.\d+)?)\s*字符", rule_text)
    if match:
        return float(match.group(1)) * 12.0
    match = re.search(r"首行缩进\s*(\d+(?:\.\d+)?)\s*(pt|磅)", rule_text, re.IGNORECASE)
    if match:
        return float(match.group(1))
    return None


def check_page_setup_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    sections = docx_evidence.get("sections") or []
    if not sections:
        return issues
    for rule in rules:
        if not rule_mentions_any(rule, ["页边距", "纸张", "a4", "margin", "页眉", "页脚"]):
            continue
        rule_text = json.dumps(rule, ensure_ascii=False)
        expected_mm = parse_expected_mm_values(rule_text)
        for section in sections:
            section_index = section.get("index")
            if "a4" in rule_text.lower() or "A4" in rule_text:
                width = section.get("page_width_mm")
                height = section.get("page_height_mm")
                if width and height and not (same_number(width, 210.0, 2.0) and same_number(height, 297.0, 2.0)):
                    issues.append(
                        build_python_issue(
                            rule,
                            "page_setup",
                            f"程序读取到纸张尺寸为 {width}mm x {height}mm",
                            "A4 纸张约为 210mm x 297mm",
                            "请将页面纸张设置为 A4。",
                            {"section": section_index, "object_type": "page"},
                        )
                    )
            for key, expected in expected_mm.items():
                actual = section.get(key)
                if actual is not None and not same_number(actual, expected, 1.0):
                    issues.append(
                        build_python_issue(
                            rule,
                            "page_setup",
                            f"程序读取到 {key} 为 {actual}mm",
                            f"规范要求为 {expected}mm",
                            "请按规范调整页面设置。",
                            {"section": section_index, "object_type": "page"},
                        )
                    )
    return issues


def check_paragraph_spacing_and_indent_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    paragraphs = docx_evidence.get("paragraphs") or []
    for rule in rules:
        if not rule_mentions_any(rule, ["行距", "段前", "段后", "缩进", "spacing", "indent"]):
            continue
        rule_text = json.dumps(rule, ensure_ascii=False)
        expected_line = parse_expected_line_value(rule_text)
        expected_spacing = parse_expected_spacing_points(rule_text)
        expected_indent = parse_expected_first_line_indent_pt(rule_text)
        for paragraph in paragraphs_matching_rule_scope(rule, paragraphs):
            paragraph_props = paragraph.get("effective_paragraph_properties") or {}
            spacing = paragraph_props.get("spacing") or {}
            indent = paragraph_props.get("indent") or {}
            if expected_line is not None and spacing.get("line") is not None:
                try:
                    actual_line = int(spacing.get("line"))
                except (TypeError, ValueError):
                    actual_line = None
                if actual_line is not None and actual_line != expected_line:
                    issues.append(
                        build_python_issue(
                            rule,
                            "paragraph_format",
                            f"程序读取到行距值为 {actual_line}",
                            f"规范要求行距值为 {expected_line}",
                            "请按规范调整该段落行距。",
                            {"paragraph_index": paragraph.get("index"), "object_type": "paragraph", "quote": str(paragraph.get("text") or "")[:60]},
                            str(paragraph.get("text") or "")[:60],
                        )
                    )
            for key, expected in expected_spacing.items():
                actual = spacing.get(key)
                if actual is not None and not same_number(actual, expected, 0.5):
                    issues.append(
                        build_python_issue(
                            rule,
                            "paragraph_format",
                            f"程序读取到 {key} 为 {format_pt(actual)}",
                            f"规范要求为 {format_pt(expected)}",
                            "请按规范调整段前段后间距。",
                            {"paragraph_index": paragraph.get("index"), "object_type": "paragraph", "quote": str(paragraph.get("text") or "")[:60]},
                            str(paragraph.get("text") or "")[:60],
                        )
                    )
            actual_indent = indent.get("first_line_pt")
            if expected_indent is not None and actual_indent is not None and not same_number(actual_indent, expected_indent, 1.5):
                issues.append(
                    build_python_issue(
                        rule,
                        "paragraph_format",
                        f"程序读取到首行缩进为 {format_pt(actual_indent)}",
                        f"规范要求首行缩进约为 {format_pt(expected_indent)}",
                        "请按规范调整首行缩进。",
                        {"paragraph_index": paragraph.get("index"), "object_type": "paragraph", "quote": str(paragraph.get("text") or "")[:60]},
                        str(paragraph.get("text") or "")[:60],
                    )
                )
    return issues


def check_table_format_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    tables = docx_evidence.get("tables") or []
    paragraphs = docx_evidence.get("paragraphs") or []
    table_captions = [p for p in paragraphs if "table_caption" in set(p.get("content_roles") or [])]
    for rule in rules:
        if not rule_mentions_any(rule, ["表格", "表题", "三线表", "table"]):
            continue
        if rule_mentions_any(rule, ["表题", "编号", "表格标题"]):
            for caption in table_captions:
                text = str(caption.get("text") or "")
                props = caption.get("effective_paragraph_properties") or {}
                if "居中" in json.dumps(rule, ensure_ascii=False) and (props.get("alignment") or "left") != "center":
                    issues.append(build_python_issue(rule, "table_format", "表题未居中", "表题应居中", "请将表题设置为居中。", {"paragraph_index": caption.get("index"), "object_type": "table_caption", "quote": text[:60]}, text[:60]))
                if re.search(r"表\s*\d+(?:[-.]\d+)?", text) is None:
                    issues.append(build_python_issue(rule, "table_format", "表题编号格式不符合常见要求", "表题编号应类似 表2-1", "请按规范修改表题编号。", {"paragraph_index": caption.get("index"), "object_type": "table_caption", "quote": text[:60]}, text[:60]))
        if rule_mentions_any(rule, ["三线表", "边线", "边框"]):
            for table in tables:
                borders = table.get("borders") or {}
                has_vertical = any((borders.get(name) or {}).get("val") not in {None, "nil", "none"} for name in ("left", "right", "insideV"))
                missing_horizontal = not borders.get("top") or not borders.get("bottom")
                if has_vertical or missing_horizontal:
                    issues.append(
                        build_python_issue(
                            rule,
                            "table_format",
                            "程序读取到表格边框不符合三线表特征",
                            "三线表通常应有顶线、底线和必要内部横线，避免竖线",
                            "请按三线表规范调整表格边框。",
                            {"object_type": "table", "section": table.get("index")},
                        )
                    )
    return issues


def check_figure_format_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    images = docx_evidence.get("images") or []
    paragraphs = docx_evidence.get("paragraphs") or []
    figure_captions = [p for p in paragraphs if "figure_caption" in set(p.get("content_roles") or [])]
    for rule in rules:
        if not rule_mentions_any(rule, ["图片", "插图", "图题", "figure", "image"]):
            continue
        if images and len(figure_captions) < len(images):
            issues.append(build_python_issue(rule, "figure_format", "图片数量多于可识别图题数量", "每张图片应有图题", "请为图片补充规范图题。", {"object_type": "figure"}))
        for caption in figure_captions:
            text = str(caption.get("text") or "")
            props = caption.get("effective_paragraph_properties") or {}
            if "居中" in json.dumps(rule, ensure_ascii=False) and (props.get("alignment") or "left") != "center":
                issues.append(build_python_issue(rule, "figure_format", "图题未居中", "图题应居中", "请将图题设置为居中。", {"paragraph_index": caption.get("index"), "object_type": "figure_caption", "quote": text[:60]}, text[:60]))
            if re.search(r"图\s*\d+(?:[-.]\d+)?", text) is None:
                issues.append(build_python_issue(rule, "figure_format", "图题编号格式不符合常见要求", "图题编号应类似 图2-1", "请按规范修改图题编号。", {"paragraph_index": caption.get("index"), "object_type": "figure_caption", "quote": text[:60]}, text[:60]))
        if rule_mentions_any(rule, ["清晰度", "dpi", "分辨率"]):
            issues.append(build_python_issue(rule, "figure_format", "DOCX 静态证据未提供可靠图片 DPI", "图片清晰度应满足规范", "请人工确认图片分辨率或使用渲染/图像元数据检查。", {"object_type": "figure"}, manual=True))
    return issues


def check_header_footer_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    headers_footers = docx_evidence.get("headers_footers") or {}
    headers = headers_footers.get("headers") or []
    footers = headers_footers.get("footers") or []
    for rule in rules:
        if not rule_mentions_any(rule, ["页眉", "页脚", "页码", "header", "footer", "page number"]):
            continue
        wants_header_content = rule_mentions_any(rule, ["页眉内容", "页眉应有", "奇数页", "偶数页", "论文题目", "章标题", "header content"])
        wants_header_font = rule_mentions_any(rule, ["页眉字体", "页眉用", "header font"])
        wants_footer_page = rule_mentions_any(rule, ["页码", "page number"])
        wants_footer_content = rule_mentions_any(rule, ["页脚内容", "footer content"])
        if wants_header_content:
            if not headers:
                issues.append(build_python_issue(rule, "header_footer", "文档中未读取到页眉", "规范要求存在页眉", "请按规范设置页眉。", {"object_type": "header"}))
            for header in headers:
                if not str(header.get("text") or "").strip():
                    issues.append(build_python_issue(rule, "header_footer", "页眉内容为空", "页眉应包含规范要求的内容", "请补充页眉内容。", {"object_type": "header", "quote": ""}))
        if wants_header_font:
            issues.append(build_python_issue(rule, "header_footer", "当前页眉解析只读取到文本和域，未稳定读取页眉字体字号", "页眉字体字号应符合规范", "请人工确认页眉字体字号，或后续扩展页眉段落格式解析。", {"object_type": "header"}, manual=True))
        if wants_footer_page or wants_footer_content:
            if not footers:
                issues.append(build_python_issue(rule, "header_footer", "文档中未读取到页脚", "规范要求存在页脚或页码", "请按规范设置页脚页码。", {"object_type": "footer"}))
                continue
            for footer in footers:
                fields = footer.get("fields") or []
                has_page_field = any("PAGE" in str(field.get("instruction") or "").upper() for field in fields if isinstance(field, dict))
                if wants_footer_page and not has_page_field:
                    issues.append(build_python_issue(rule, "header_footer", "页脚中未读取到 PAGE 页码域", "页脚应包含页码", "请插入规范要求的页码。", {"object_type": "footer", "quote": str(footer.get("text") or "")[:60]}))
    return issues


def extract_reference_numbers(paragraphs: list[dict[str, Any]]) -> set[int]:
    refs: set[int] = set()
    for paragraph in paragraphs:
        if "reference" not in set(paragraph.get("content_roles") or []):
            continue
        match = re.match(r"\s*\[(\d+)\]", str(paragraph.get("text") or ""))
        if match:
            refs.add(int(match.group(1)))
    return refs


def extract_body_citation_numbers(paragraphs: list[dict[str, Any]]) -> list[tuple[int, int, str]]:
    citations: list[tuple[int, int, str]] = []
    for paragraph in paragraphs:
        if "reference" in set(paragraph.get("content_roles") or []):
            continue
        text = str(paragraph.get("text") or "")
        for citation in re.findall(r"\[(\d+(?:\s*[-,，]\s*\d+)*)\]", text):
            for number in re.findall(r"\d+", citation):
                citations.append((int(number), int(paragraph.get("index") or 0), text[:60]))
    return citations


def check_reference_format_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    paragraphs = docx_evidence.get("paragraphs") or []
    refs = extract_reference_numbers(paragraphs)
    citations = extract_body_citation_numbers(paragraphs)
    for rule in rules:
        if not rule_mentions_any(rule, ["参考文献", "引用", "citation", "reference"]):
            continue
        for number, paragraph_index, quote in citations:
            if refs and number not in refs:
                issues.append(build_python_issue(rule, "citation_reference", f"正文引用 [{number}] 未在参考文献列表中找到", "正文引用编号应能对应参考文献列表", "请补充对应参考文献或修改引用编号。", {"paragraph_index": paragraph_index, "object_type": "citation", "quote": quote}, quote))
        for paragraph in paragraphs:
            if "reference" in set(paragraph.get("content_roles") or []) and not re.match(r"\s*\[\d+\]", str(paragraph.get("text") or "")):
                text = str(paragraph.get("text") or "")
                issues.append(build_python_issue(rule, "citation_reference", "参考文献列表项缺少 [编号] 格式", "参考文献应以 [编号] 开头", "请按规范补充参考文献编号。", {"paragraph_index": paragraph.get("index"), "object_type": "reference", "quote": text[:60]}, text[:60]))
    return issues


def check_formula_structural_rules(rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    formulas = [p for p in docx_evidence.get("paragraphs") or [] if "formula" in set(p.get("content_roles") or [])]
    for rule in rules:
        if not rule_mentions_any(rule, ["公式", "formula", "equation"]):
            continue
        rule_text = json.dumps(rule, ensure_ascii=False)
        for paragraph in formulas:
            text = str(paragraph.get("text") or "")
            props = paragraph.get("effective_paragraph_properties") or {}
            if "居中" in rule_text and (props.get("alignment") or "left") != "center":
                issues.append(build_python_issue(rule, "formula_format", "公式段落未居中", "公式应居中编排", "请将公式段落设置为居中。", {"paragraph_index": paragraph.get("index"), "object_type": "formula", "quote": text[:60]}, text[:60]))
            has_any_formula_number = re.search(r"[\(\uff08]\s*\d+(?:[-.]\d+)?\s*[\)\uff09]", text)
            has_chapter_formula_number = re.search(r"[\(\uff08]\s*\d+[-.]\d+\s*[\)\uff09]", text)
            if rule_mentions_any(rule, ["编号", "number"]) and has_any_formula_number and not has_chapter_formula_number:
                issues.append(build_python_issue(rule, "formula_format", "公式编号格式不符合章号-序号格式", "公式编号应类似 (3-1)", "请按规范修改公式编号。", {"paragraph_index": paragraph.get("index"), "object_type": "formula", "quote": text[:60]}, text[:60]))
    return issues


def run_structural_format_checks(normalized_rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    python_rules = filter_rules_by_check_method(normalized_rules, "python")
    issues: list[dict[str, Any]] = []
    issues.extend(check_page_setup_rules(python_rules, docx_evidence))
    issues.extend(check_paragraph_spacing_and_indent_rules(python_rules, docx_evidence))
    issues.extend(check_table_format_rules(python_rules, docx_evidence))
    issues.extend(check_figure_format_rules(python_rules, docx_evidence))
    issues.extend(check_header_footer_rules(python_rules, docx_evidence))
    issues.extend(check_reference_format_rules(python_rules, docx_evidence))
    issues.extend(check_formula_structural_rules(python_rules, docx_evidence))
    return issues


def run_deterministic_checks(normalized_rules: list[dict[str, Any]], docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    paragraphs = docx_evidence.get("paragraphs") or []
    for rule in normalized_rules:
        if normalize_check_method(rule.get("check_method")) != "python":
            continue
        rule_text = json.dumps(rule, ensure_ascii=False)
        expected_font = next(
            (
                font
                for font in ["\u5b8b\u4f53", "SimSun", "\u9ed1\u4f53", "\u6977\u4f53", "Times New Roman", "Arial", "Calibri"]
                if font in rule_text
            ),
            None,
        )
        expected_font_size = parse_expected_font_size_pt(rule_text)
        expected_alignment = parse_expected_alignment(rule_text)
        if not expected_font and expected_font_size is None and not expected_alignment:
            continue
        rule_script = expected_font_script(rule_text, expected_font)
        for paragraph in paragraphs_matching_rule_scope(rule, paragraphs):
            if not isinstance(paragraph, dict):
                continue
            paragraph_props = paragraph.get("effective_paragraph_properties") or {}
            if expected_alignment:
                actual_alignment = paragraph_props.get("alignment") or "left"
                if actual_alignment != expected_alignment:
                    issues.append(build_deterministic_issue(rule, paragraph, str(actual_alignment), expected_alignment))
                    if len(issues) >= 200:
                        return issues
                    continue
            runs = paragraph.get("runs") or []
            for run in runs:
                if expected_font:
                    actual_font, should_check = select_run_font_for_rule(run, rule_script)
                    if not should_check:
                        continue
                    if actual_font and expected_font not in str(actual_font):
                        issues.append(build_deterministic_issue(rule, paragraph, str(actual_font), expected_font))
                        break
                if expected_font_size is not None and str(run.get("text") or "").strip():
                    actual_size = (run.get("effective_properties") or {}).get("font_size_pt")
                    if actual_size is not None and not same_number(actual_size, expected_font_size):
                        issues.append(
                            build_deterministic_issue(
                                rule,
                                paragraph,
                                format_pt(actual_size),
                                format_pt(expected_font_size),
                            )
                        )
                        break
            if len(issues) >= 200:
                return issues
    return issues


def rule_mentions_any(rule: dict[str, Any], keywords: list[str]) -> bool:
    text = json.dumps(rule, ensure_ascii=False).lower()
    return any(keyword.lower() in text for keyword in keywords)


IMPLEMENTED_PYTHON_CHECKS = [
    "font_family",
    "font_size",
    "paragraph_alignment",
    "paragraph_spacing",
    "paragraph_indent",
    "page_setup",
    "header",
    "footer_page_number",
    "table_format",
    "figure_format",
    "citation_reference",
    "formula_missing_number",
    "toc_rendered_page_consistency",
]


def infer_required_checker_names(rule: dict[str, Any]) -> set[str]:
    text = json.dumps(rule, ensure_ascii=False).lower()
    required: set[str] = set()
    if any(keyword in text for keyword in ["font", "字体", "times new roman", "宋体", "黑体", "楷体", "arial", "calibri"]):
        required.add("font_family")
    if any(keyword in text for keyword in ["字号", "font size", "pt", "磅", "小四", "小二", "三号", "四号", "五号"]):
        required.add("font_size")
    if any(keyword in text for keyword in ["对齐", "居中", "居左", "居右", "alignment", "center", "left", "right", "justify"]):
        required.add("paragraph_alignment")
    if any(keyword in text for keyword in ["行距", "段前", "段后", "line spacing", "spacing"]):
        required.add("paragraph_spacing")
    if any(keyword in text for keyword in ["缩进", "indent", "首行"]):
        required.add("paragraph_indent")
    if any(keyword in text for keyword in ["页边距", "纸张", "a4", "margin", "page setup"]):
        required.add("page_setup")
    if any(keyword in text for keyword in ["页眉", "header"]):
        required.add("header")
    if any(keyword in text for keyword in ["页脚", "页码", "footer", "page number"]):
        required.add("footer_page_number")
    if any(keyword in text for keyword in ["表格", "三线表", "表题", "table"]):
        required.add("table_format")
    if any(keyword in text for keyword in ["图片", "插图", "图题", "figure", "image"]):
        required.add("figure_format")
    if any(keyword in text for keyword in ["目录", "toc"]):
        required.add("toc_rendered_page_consistency")
    if any(keyword in text for keyword in ["公式", "formula", "equation"]):
        required.add("formula_missing_number")
    if any(keyword in text for keyword in ["参考文献", "引用", "citation", "reference"]):
        required.add("citation_reference")
    return required


def rule_has_dedicated_python_checker(rule: dict[str, Any]) -> bool:
    required = infer_required_checker_names(rule)
    if not required:
        return False
    return required.issubset(set(IMPLEMENTED_PYTHON_CHECKS))


def build_check_coverage_audit(format_rule: dict[str, Any], docx_evidence: dict[str, Any]) -> dict[str, Any]:
    normalized_rules = normalize_format_rules(format_rule)
    python_rules = filter_rules_by_check_method(normalized_rules, "python")
    llm_rules = filter_rules_by_check_method(normalized_rules, "llm")
    role_counts: Counter[str] = Counter()
    for paragraph in docx_evidence.get("paragraphs") or []:
        for role in paragraph.get("content_roles") or classify_paragraph_content_roles(paragraph):
            role_counts[role] += 1

    unsupported_rules: list[dict[str, Any]] = []
    required_checkers: Counter[str] = Counter()
    for rule in python_rules:
        required = infer_required_checker_names(rule)
        for checker in required:
            required_checkers[checker] += 1
        if not rule_has_dedicated_python_checker(rule):
            unsupported_rules.append(
                {
                    "rule_id": rule.get("rule_id") or rule.get("id"),
                    "category": rule.get("category"),
                    "scope": rule.get("scope"),
                    "rule_name": rule.get("rule_name") or rule.get("description"),
                    "required_checkers": sorted(required),
                }
            )

    areas_requiring_more_checkers = sorted(
        {
            checker.split("_", 1)[0]
            for rule in unsupported_rules
            for checker in rule.get("required_checkers", [])
            if checker not in IMPLEMENTED_PYTHON_CHECKS
        }
    )
    return {
        "implemented_python_checks": IMPLEMENTED_PYTHON_CHECKS,
        "known_unimplemented_python_checks": [
            "odd_even_rendered_header_validation",
            "cover_page_rendered_no_page_number_validation",
            "cross_page_table_repeat_header_validation",
            "reliable_image_dpi_validation",
            "full_reference_style_semantic_classification",
        ],
        "evidence_summary": {
            "paragraph_count": docx_evidence.get("paragraph_count", len(docx_evidence.get("paragraphs") or [])),
            "section_count": docx_evidence.get("section_count"),
            "table_count": docx_evidence.get("table_count", len(docx_evidence.get("tables") or [])),
            "image_count": len(docx_evidence.get("images") or []),
            "header_count": len((docx_evidence.get("headers_footers") or {}).get("headers") or []),
            "footer_count": len((docx_evidence.get("headers_footers") or {}).get("footers") or []),
            "footnote_count": len(docx_evidence.get("footnotes") or []),
            "endnote_count": len(docx_evidence.get("endnotes") or []),
        },
        "paragraph_role_counts": dict(sorted(role_counts.items())),
        "python_rule_count": len(python_rules),
        "llm_rule_count": len(llm_rules),
        "required_checker_counts": dict(sorted(required_checkers.items())),
        "python_rules_without_dedicated_checker": unsupported_rules,
        "areas_requiring_more_checkers": areas_requiring_more_checkers,
    }


def extract_formulas_from_docx(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() != ".docx" or not path.exists():
        return []
    try:
        parts = read_docx_parts(path)
        document_root = ET.fromstring(parts["word/document.xml"])
    except Exception:
        return []

    formulas: list[dict[str, Any]] = []
    paragraphs = list(document_root.iter(qn(W_NS, "p")))
    formula_tags = {qn(M_NS, "oMath"), qn(M_NS, "oMathPara")}
    for paragraph_index, paragraph in enumerate(paragraphs, start=1):
        math_nodes = [node for node in paragraph.iter() if node.tag in formula_tags]
        if not math_nodes:
            continue
        text = paragraph_text(paragraph)
        formulas.append(
            {
                "paragraph_index": paragraph_index,
                "text": text,
                "formula_count": len(math_nodes),
                "has_formula_number": bool(re.search(r"[\(\uff08]\s*\d+(?:[.-]\d+)?\s*[\)\uff09]", text)),
                "alignment": element_attr(child(child(paragraph, W_NS, "pPr"), W_NS, "jc"), W_NS, "val"),
                "raw_math_xml": [ET.tostring(node, encoding="unicode") for node in math_nodes],
            }
        )
    return formulas


def check_formula_format_from_docx(path: Path, normalized_rules: list[dict[str, Any]]) -> list[dict[str, Any]]:
    formula_rules = [
        rule
        for rule in normalized_rules
        if normalize_check_method(rule.get("check_method")) == "python"
        and rule_mentions_any(rule, ["formula", "\u516c\u5f0f"])
    ]
    if not formula_rules:
        return []

    issues: list[dict[str, Any]] = []
    formulas = extract_formulas_from_docx(path)
    for rule in formula_rules:
        rule_id = rule.get("rule_id") or rule.get("id") or "R-FORMULA"
        if not rule_mentions_any(rule, ["number", "\u7f16\u53f7"]):
            continue
        for formula in formulas:
            if formula.get("has_formula_number"):
                continue
            quote = str(formula.get("text") or "\u516c\u5f0f")[:60]
            issues.append(
                normalize_issue(
                    {
                        "issue_id": f"D-{rule_id}-{formula.get('paragraph_index')}",
                        "rule_id": rule_id,
                        "category": rule.get("category") or "formula_format",
                        "severity": rule.get("severity") or "medium",
                        "reason": "\u7a0b\u5e8f\u8bfb\u53d6\u5230\u516c\u5f0f\u6bb5\u843d\uff0c\u4f46\u672a\u68c0\u6d4b\u5230\u516c\u5f0f\u7f16\u53f7",
                        "expected": rule.get("description") or rule.get("expected") or "\u516c\u5f0f\u5e94\u6309\u89c4\u8303\u7f16\u53f7",
                        "suggestion": "\u8bf7\u4e3a\u8be5\u516c\u5f0f\u8865\u5145\u89c4\u8303\u8981\u6c42\u7684\u7f16\u53f7",
                        "need_manual_confirmation": False,
                        "source": "python_hard_checker",
                        "location": {
                            "page": None,
                            "section": None,
                            "paragraph_index": formula.get("paragraph_index"),
                            "object_type": "formula",
                            "quote": quote,
                        },
                        "anchor_text": quote,
                    },
                    1,
                )
            )
    return issues


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path | None:
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(docx_path),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
        )
    except (FileNotFoundError, subprocess.CalledProcessError):
        return None
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    return pdf_path if pdf_path.exists() else None


def extract_pdf_pages(pdf_path: Path) -> list[dict[str, Any]]:
    try:
        import fitz  # type: ignore
    except ImportError:
        return []
    pages: list[dict[str, Any]] = []
    try:
        document = fitz.open(pdf_path)
        for index, page in enumerate(document, start=1):
            pages.append({"page": index, "text": page.get_text("text")})
    except Exception:
        return []
    return pages


def extract_toc_entries_from_text(text: str) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for line in text.splitlines():
        stripped = line.strip()
        match = re.match(r"^(.{2,80}?)[.\s·…]{2,}(\d+)$", stripped)
        if not match:
            continue
        entries.append({"title": match.group(1).strip(), "page": int(match.group(2)), "raw": stripped})
    return entries


def extract_heading_candidates(docx_evidence: dict[str, Any]) -> list[dict[str, Any]]:
    headings: list[dict[str, Any]] = []
    for paragraph in docx_evidence.get("paragraphs") or []:
        if not isinstance(paragraph, dict):
            continue
        style_id = str(paragraph.get("style_id") or "")
        text = str(paragraph.get("text") or "").strip()
        if not text:
            continue
        if style_id.lower().startswith("heading") or re.match(r"^\s*(\u7b2c.+\u7ae0|\d+(?:\.\d+)*)", text):
            headings.append({"text": text, "paragraph_index": paragraph.get("index")})
    return headings


def check_toc_page_consistency_from_docx(path: Path, docx_evidence: dict[str, Any], normalized_rules: list[dict[str, Any]]) -> list[dict[str, Any]]:
    page_rules = [
        rule
        for rule in normalized_rules
        if normalize_check_method(rule.get("check_method")) == "python"
        and rule_mentions_any(rule, ["toc", "\u76ee\u5f55", "\u9875\u7801"])
    ]
    if not page_rules:
        return []
    toc_entries = extract_toc_entries_from_text(str(docx_evidence.get("full_text") or ""))
    headings = extract_heading_candidates(docx_evidence)
    if not toc_entries or not headings:
        return []

    render_dir = path.parent / ".rendered_pages"
    pdf_path = convert_docx_to_pdf(path, render_dir)
    if pdf_path is None:
        return []
    pages = extract_pdf_pages(pdf_path)
    if not pages:
        return []

    actual_pages: dict[str, int] = {}
    for heading in headings:
        title = heading["text"]
        for page in pages:
            if title in page.get("text", ""):
                actual_pages[title] = int(page["page"])
                break

    issues: list[dict[str, Any]] = []
    rule = page_rules[0]
    rule_id = rule.get("rule_id") or rule.get("id") or "R-PAGE"
    for entry in toc_entries:
        actual = actual_pages.get(entry["title"])
        if actual is None or actual == entry["page"]:
            continue
        issues.append(
            normalize_issue(
                {
                    "issue_id": f"D-{rule_id}-{entry['page']}",
                    "rule_id": rule_id,
                    "category": "page_numbering",
                    "severity": rule.get("severity") or "high",
                    "reason": "\u76ee\u5f55\u9875\u7801\u4e0e\u6e32\u67d3\u540e\u6b63\u6587\u5b9e\u9645\u9875\u7801\u4e0d\u4e00\u81f4",
                    "expected": f"\u76ee\u5f55\u9875\u7801\uff1a{entry['page']}",
                    "suggestion": f"\u8bf7\u66f4\u65b0\u76ee\u5f55\uff0c\u8be5\u6807\u9898\u6e32\u67d3\u540e\u5b9e\u9645\u5728\u7b2c {actual} \u9875",
                    "need_manual_confirmation": False,
                    "source": "python_hard_checker",
                    "location": {
                        "page": entry["page"],
                        "section": entry["title"],
                        "paragraph_index": None,
                        "object_type": "page",
                        "quote": entry["raw"],
                    },
                    "anchor_text": entry["title"],
                },
                1,
            )
        )
    return issues


def merge_issues(primary: list[dict[str, Any]], secondary: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[tuple[str, str, str]] = set()
    merged: list[dict[str, Any]] = []
    for issue in primary + secondary:
        location = issue.get("location") if isinstance(issue.get("location"), dict) else {}
        key = (
            str(issue.get("rule_id") or issue.get("category") or ""),
            str(location.get("paragraph_index") or ""),
            str(issue.get("anchor_text") or location.get("quote") or "")[:80],
        )
        if key in seen:
            continue
        seen.add(key)
        merged.append(issue)
    return merged


def enrich_hybrid_result(
    format_rule: dict[str, Any],
    check_result: dict[str, Any],
    checked_file: Path,
    workflow: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    docx_evidence = build_docx_analysis_evidence(checked_file)
    normalized_rules = normalize_format_rules(format_rule)
    deterministic_issues = run_deterministic_checks(normalized_rules, docx_evidence)
    llm_issues = [
        normalize_issue(issue, index)
        for index, issue in enumerate(check_result.get("issues", []), start=1)
        if isinstance(issue, dict)
    ]
    issues = merge_issues(deterministic_issues, llm_issues)
    statistics = recompute_statistics(issues, len(normalized_rules) or len(format_rule.get("rules", [])))
    format_rule["normalized_rules"] = normalized_rules
    check_result["issues"] = issues
    check_result["llm_raw"] = {
        **(check_result.get("llm_raw") or {}),
        "normalized_rule_count": len(normalized_rules),
        "deterministic_issue_count": len(deterministic_issues),
        "semantic_issue_count": len(llm_issues),
        "docx_evidence_summary": {
            "parser": docx_evidence.get("parser"),
            "paragraph_count": docx_evidence.get("paragraph_count"),
            "section_count": docx_evidence.get("section_count"),
            "table_count": docx_evidence.get("table_count"),
        },
    }
    return set_result_workflow(format_rule, check_result, statistics, workflow)


def build_compare_result(llm_result: dict[str, Any], hybrid_result: dict[str, Any]) -> dict[str, Any]:
    llm_count = len(llm_result.get("issues", []))
    hybrid_count = len(hybrid_result.get("issues", []))
    return {
        "llm_direct_total_issues": llm_count,
        "hybrid_total_issues": hybrid_count,
        "deterministic_added_issues": max(hybrid_count - llm_count, 0),
        "recommended_workflow": "hybrid",
        "reason": "hybrid workflow combines python-docx/XML deterministic evidence with LLM semantic judgment.",
    }


def build_placeholder_outputs(
    standard_file: Path,
    checked_file: Path,
    format_rule_output: Path,
    workflow: str = "llm_direct",
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule = {
        "file_analysis": {
            "file_name": standard_file.name,
            "file_type": standard_file.suffix.lstrip(".").lower(),
            "selected_parser": "placeholder-parser",
            "analysis_summary": "占位规则文件，真实规则抽取由 LLM 完成。",
        },
        "rules": [],
    }
    statistics = {
        "total_issues": 0,
        "by_category": {},
        "by_severity": {},
        "need_manual_confirmation": 0,
        "checked_rule_count": 0,
    }
    check_result = {
        "paper_analysis": {
            "standard_file": format_rule_output.name,
            "paper_file": checked_file.name,
            "paper_type": "docx",
            "selected_parser": "placeholder-docx-parser",
            "check_summary": "占位检测结果，真实格式检查由规则引擎和 LLM 完成。",
        },
        "statistics": statistics,
        "issues": [],
    }
    if workflow == "compare":
        check_result["compare_result"] = build_compare_result(check_result, check_result)
    return set_result_workflow(format_rule, check_result, statistics, workflow)


def extract_format_rules_with_llm(
    settings: dict[str, Any],
    standard_file: Path,
    llm_log_output: Path,
) -> dict[str, Any]:
    standard_text = read_document_text(standard_file)
    extract_rule_prompt = force_chinese_output(
        read_skill_prompt(
            EXTRACT_RULE_PROMPT_PATH,
            "Extract format rules from the standard document and return JSON with rules.",
        )
    )
    rule_content = call_openai_compatible_chat(
        settings,
        [
            {"role": "system", "content": extract_rule_prompt},
            {
                "role": "user",
                "content": (
                    f"请从以下标准格式规范文件中抽取可执行的规则，并输出 JSON。\n\n"
                    f"规范文件名：{standard_file.name}\n\n"
                    f"规范全文：\n{standard_text}"
                ),
            },
        ],
        agent_name="agent1_extract_format_rules",
        log_path=llm_log_output,
    )
    rule_payload = parse_llm_json(rule_content, "agent1_extract_format_rules")
    rules = normalize_format_rules({"rules": rule_payload.get("rules", [])})
    return {
        "file_analysis": {
            "file_name": standard_file.name,
            "file_type": standard_file.suffix.lstrip(".").lower(),
            "selected_parser": "llm-document-reader",
            "llm_provider": settings["provider"],
            "llm_model": settings["model"],
            "analysis_summary": "已由 LLM 根据规范文件抽取格式规则，并按 python/llm 检查方式归类。",
        },
        "rules": rules,
        "llm_raw": rule_payload,
    }


def run_hard_format_from_rule(
    format_rule: dict[str, Any],
    checked_file: Path,
    workflow: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    docx_evidence = build_docx_analysis_evidence(checked_file)
    normalized_rules = normalize_format_rules(format_rule)
    python_rules = filter_rules_by_check_method(normalized_rules, "python")
    coverage_audit = build_check_coverage_audit(format_rule, docx_evidence)
    issues = []
    issues.extend(run_deterministic_checks(python_rules, docx_evidence))
    issues.extend(run_structural_format_checks(python_rules, docx_evidence))
    issues.extend(check_formula_format_from_docx(checked_file, python_rules))
    issues.extend(check_toc_page_consistency_from_docx(checked_file, docx_evidence, python_rules))
    issues = [normalize_issue(issue, index) for index, issue in enumerate(issues, start=1)]
    statistics = recompute_statistics(issues, len(python_rules))
    format_rule["normalized_rules"] = normalized_rules
    check_result = {
        "paper_analysis": {
            "paper_file": checked_file.name,
            "paper_type": checked_file.suffix.lstrip(".").lower() or "docx",
            "selected_parser": "python-docx+xml + Office Open XML + optional PDF rendering",
            "check_summary": "已使用 Python 检查段落、页面、页眉页脚、表格、图片题注、参考文献编号、公式和目录等硬性格式；需渲染或人工确认的项目见 coverage_audit。",
        },
        "statistics": statistics,
        "issues": issues,
        "llm_raw": {
            "normalized_rule_count": len(normalized_rules),
            "python_rule_count": len(python_rules),
            "coverage_audit": coverage_audit,
            "docx_evidence_summary": {
                "parser": docx_evidence.get("parser"),
                "paragraph_count": docx_evidence.get("paragraph_count"),
                "section_count": docx_evidence.get("section_count"),
                "table_count": docx_evidence.get("table_count"),
                "formula_count": len(extract_formulas_from_docx(checked_file)),
            },
        },
    }
    return set_result_workflow(format_rule, check_result, statistics, workflow)


def run_complex_format_explanation_from_evidence(
    settings: dict[str, Any],
    format_rule: dict[str, Any],
    checked_file: Path,
    hard_result: dict[str, Any],
    llm_log_output: Path,
    workflow: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    checked_text = read_document_text(checked_file)
    docx_evidence = build_docx_analysis_evidence(checked_file)
    normalized_rules = normalize_format_rules(format_rule)
    python_rules = filter_rules_by_check_method(normalized_rules, "python")
    prompt = force_chinese_output(
        "你是复杂格式规则解释 Agent。你只解释 Python、WordprocessingML、LibreOffice/PDF 渲染或图片元数据已经提取出的格式证据，"
        "判断这些复杂证据是否违反格式规范，并生成可批注的中文错误说明。"
        "你不得检查语法、错别字、病句、摘要语义一致性、术语一致性等语言语义问题；这些由 agent3_language_semantic_checker 处理。"
        "你可以处理奇偶页页眉、封面实际无页码、跨页表格续表、图片 DPI、图题表题与格式证据的解释。"
        "请只输出合法 JSON，字段包括 paper_analysis、statistics、issues。每个 issue 必须包含 location、anchor_text、reason、expected、suggestion、comment_text。"
    )
    content = call_openai_compatible_chat(
        settings,
        [
            {"role": "system", "content": prompt},
            {
                "role": "user",
                "content": (
                    f"Python 格式类规则 JSON：\n{json.dumps(python_rules, ensure_ascii=False)}\n\n"
                    f"Python 硬格式检查结果 JSON：\n{json.dumps(hard_result, ensure_ascii=False)}\n\n"
                    f"python-docx 结构化证据摘要 JSON：\n{json.dumps({'parser': docx_evidence.get('parser'), 'paragraph_count': docx_evidence.get('paragraph_count'), 'section_count': docx_evidence.get('section_count'), 'table_count': docx_evidence.get('table_count'), 'headers_footers': docx_evidence.get('headers_footers'), 'images': docx_evidence.get('images')}, ensure_ascii=False)}\n\n"
                    "任务：只基于上述格式证据解释复杂格式问题，不要检查语言语义。\n\n"
                    f"待检测论文文件名：{checked_file.name}\n\n"
                    f"论文全文定位辅助：\n{checked_text[:12000]}"
                ),
            },
        ],
        agent_name="agent2_complex_format_explainer",
        log_path=llm_log_output,
    )
    payload = parse_llm_json(content, "agent2_complex_format_explainer")
    issues = [
        normalize_issue({**issue, "source": issue.get("source") or "llm_complex_format_explainer"}, index)
        for index, issue in enumerate(payload.get("issues", []), start=1)
        if isinstance(issue, dict)
    ]
    statistics = recompute_statistics(issues, len(python_rules))
    check_result = {
        "paper_analysis": {
            **(payload.get("paper_analysis") if isinstance(payload.get("paper_analysis"), dict) else {}),
            "paper_file": checked_file.name,
            "paper_type": checked_file.suffix.lstrip(".").lower() or "docx",
            "selected_parser": "python evidence + LLM complex format explainer",
            "check_summary": "已由复杂格式解释 Agent 基于 Python 格式证据解释复杂格式问题。",
        },
        "statistics": statistics,
        "issues": issues,
        "llm_raw": payload,
    }
    return set_result_workflow(format_rule, check_result, statistics, workflow)


def run_language_semantic_check_from_rule(
    settings: dict[str, Any],
    format_rule: dict[str, Any],
    checked_file: Path,
    llm_log_output: Path,
    workflow: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    checked_text = read_document_text(checked_file)
    docx_evidence = build_docx_analysis_evidence(checked_file)
    normalized_rules = normalize_format_rules(format_rule)
    llm_rules = filter_rules_by_check_method(normalized_rules, "llm")
    prompt = force_chinese_output(
        "你是语言语义拓展检查 Agent。你只检查语法、错别字、病句、语义一致性、中英文摘要对应关系、正文引用与参考文献语义关系、术语一致性。"
        "不要检查字体、字号、页边距、页码、行距、表格边框、公式编号、图题表题格式、页眉页脚等硬性格式；这些由基础格式功能处理。"
        "请只输出合法 JSON，字段包括 paper_analysis、statistics、issues。每个 issue 必须包含 location、anchor_text、reason、expected、suggestion、comment_text。"
        "location.paragraph_index 必须尽量使用 python-docx 结构化证据中的 paragraphs[].index。"
    )
    content = call_openai_compatible_chat(
        settings,
        [
            {"role": "system", "content": prompt},
            {
                "role": "user",
                "content": (
                    f"LLM 语义类规则 JSON：\n{json.dumps(llm_rules, ensure_ascii=False)}\n\n"
                    f"全文读取概况 JSON：\n{json.dumps(build_document_overview(checked_text, [checked_text]), ensure_ascii=False)}\n\n"
                    f"python-docx 结构化证据摘要 JSON：\n{json.dumps({'parser': docx_evidence.get('parser'), 'paragraph_count': docx_evidence.get('paragraph_count'), 'section_count': docx_evidence.get('section_count'), 'table_count': docx_evidence.get('table_count')}, ensure_ascii=False)}\n\n"
                    "定位要求：每个 issue 的 location.paragraph_index 必须尽量使用 python-docx 结构化证据中的 paragraphs[].index；anchor_text 必须是论文中的连续原文短语，用于 Word 原生批注定位。\n\n"
                    f"待检测论文文件名：{checked_file.name}\n\n"
                    f"论文全文：\n{checked_text}"
                ),
            },
        ],
        agent_name="agent3_language_semantic_checker",
        log_path=llm_log_output,
    )
    payload = parse_llm_json(content, "agent3_language_semantic_checker")
    issues = [
        normalize_issue({**issue, "source": issue.get("source") or "llm_language_semantic_checker"}, index)
        for index, issue in enumerate(payload.get("issues", []), start=1)
        if isinstance(issue, dict)
    ]
    statistics = recompute_statistics(issues, len(llm_rules))
    format_rule["normalized_rules"] = normalized_rules
    check_result = {
        "paper_analysis": {
            **(payload.get("paper_analysis") if isinstance(payload.get("paper_analysis"), dict) else {}),
            "paper_file": checked_file.name,
            "paper_type": checked_file.suffix.lstrip(".").lower() or "docx",
            "selected_parser": "plain-text-reader + LLM language semantic checker",
            "check_summary": "已由语言语义拓展 Agent 检查语法、错别字、病句、语义一致性和引用对应关系。",
            "checked_text_chars": len(checked_text),
            "checked_chunk_count": 1,
            "chunk_char_lengths": [len(checked_text)],
        },
        "statistics": statistics,
        "issues": issues,
        "llm_raw": {
            "chunk_count": 1,
            "full_paper": payload,
        },
    }
    return set_result_workflow(format_rule, check_result, statistics, workflow)


def run_base_format_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule = extract_format_rules_with_llm(settings, standard_file, llm_log_output)
    return run_base_format_workflow_from_rule(settings, format_rule, checked_file, llm_log_output)


def run_base_format_workflow_from_rule(
    settings: dict[str, Any],
    format_rule: dict[str, Any],
    checked_file: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    hard_format_rule = copy.deepcopy(format_rule)
    _, hard_result, _ = run_hard_format_from_rule(hard_format_rule, checked_file, "base_format")
    _, complex_result, _ = run_complex_format_explanation_from_evidence(
        settings, format_rule, checked_file, hard_result, llm_log_output, "base_format"
    )
    issues = merge_issues(hard_result.get("issues", []), complex_result.get("issues", []))
    statistics = recompute_statistics(issues, len(normalize_format_rules(format_rule)))
    check_result = {
        "paper_analysis": {
            "paper_file": checked_file.name,
            "paper_type": checked_file.suffix.lstrip(".").lower() or "docx",
            "selected_parser": "python hard checker + LLM complex format explainer",
            "check_summary": "已完成 Python 硬格式检查和复杂格式规则解释。",
        },
        "statistics": statistics,
        "issues": issues,
        "python_hard_format_result": hard_result,
        "complex_format_result": complex_result,
        "llm_raw": {"complex_format": complex_result.get("llm_raw")},
    }
    return set_result_workflow(format_rule, check_result, statistics, "base_format")


def run_language_semantic_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule = extract_format_rules_with_llm(settings, standard_file, llm_log_output)
    return run_language_semantic_check_from_rule(settings, format_rule, checked_file, llm_log_output, "language_semantic")


def run_full_check_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule = extract_format_rules_with_llm(settings, standard_file, llm_log_output)
    base_format_rule = copy.deepcopy(format_rule)
    _, base_result, _ = run_base_format_workflow_from_rule(settings, base_format_rule, checked_file, llm_log_output)
    _, language_result, _ = run_language_semantic_check_from_rule(
        settings, format_rule, checked_file, llm_log_output, "language_semantic"
    )
    issues = merge_issues(base_result.get("issues", []), language_result.get("issues", []))
    statistics = recompute_statistics(issues, len(normalize_format_rules(format_rule)))
    check_result = {
        "paper_analysis": {
            "paper_file": checked_file.name,
            "paper_type": checked_file.suffix.lstrip(".").lower() or "docx",
            "selected_parser": "base format checker + language semantic checker",
            "check_summary": "已完成基础格式检查和语言语义拓展检查。",
            "checked_text_chars": language_result.get("paper_analysis", {}).get("checked_text_chars"),
            "checked_chunk_count": language_result.get("paper_analysis", {}).get("checked_chunk_count"),
            "chunk_char_lengths": language_result.get("paper_analysis", {}).get("chunk_char_lengths"),
        },
        "statistics": statistics,
        "issues": issues,
        "base_format_result": base_result,
        "language_semantic_result": language_result,
        "llm_raw": {
            "language_semantic": language_result.get("llm_raw"),
            "base_format": base_result.get("llm_raw"),
        },
    }
    return set_result_workflow(format_rule, check_result, statistics, "full_check")


def run_llm_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    format_rule_output: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    standard_text = read_document_text(standard_file)
    checked_text = read_document_text(checked_file)
    extract_rule_prompt = force_chinese_output(
        read_skill_prompt(
            EXTRACT_RULE_PROMPT_PATH,
            "Extract format rules from the standard document and return JSON with rules.",
        )
    )
    format_check_prompt = force_chinese_output(
        read_skill_prompt(
            FORMAT_CHECK_PROMPT_PATH,
            "Check paper format by rules and return JSON with statistics and issues.",
        )
    )

    rule_content = call_openai_compatible_chat(
        settings,
        [
            {
                "role": "system",
                "content": extract_rule_prompt,
            },
            {
                "role": "user",
                "content": f"璇蜂粠浠ヤ笅瑙勮寖鏂囦欢鍐呭涓彁鍙栬鏂囨牸寮忚鍒欍€俓n\n鏂囦欢鍚嶏細{standard_file.name}\n\n鍐呭锛歕n{standard_text}",
            },
        ],
        agent_name="agent1_extract_format_rules",
        log_path=llm_log_output,
    )
    rule_payload = parse_llm_json(rule_content, "agent1_extract_format_rules")
    format_rule = {
        "file_analysis": {
            "file_name": standard_file.name,
            "file_type": standard_file.suffix.lstrip(".").lower(),
            "selected_parser": "llm-document-reader",
            "llm_provider": settings["provider"],
            "llm_model": settings["model"],
            "analysis_summary": "已由 LLM 根据规范文件抽取格式规则。",
        },
        "rules": rule_payload.get("rules", []),
        "llm_raw": rule_payload,
    }

    docx_evidence = build_docx_analysis_evidence(checked_file)
    check_content = call_openai_compatible_chat(
        settings,
        [
            {
                "role": "system",
                "content": format_check_prompt,
            },
            {
                "role": "user",
                "content": (
                    "璇锋牴鎹牸寮忚鍒欐鏌ユ暣绡囪鏂囥€俓n"
                    "閲嶈锛氫笅闈㈢殑鈥滆鏂囧叏鏂囧唴瀹光€濇槸绯荤粺涓€娆℃€т紶鍏ョ殑瀹屾暣鏂囨湰锛屼笉瑕佹寜鍒嗙墖鐞嗚В锛屼篃涓嶈杈撳嚭鈥滆鏂囦笉瀹屾暣鈥濃€滃彧鍒版煇涓€绔犫€濃€滃彧鑳芥鏌ョ湅鍒扮殑閮ㄥ垎鈥濈瓑鍒ゆ柇銆俓n"
                    "绯荤粺宸茬粡鍦ㄦ湰鍦颁娇鐢?Python 鐨?python-docx 搴撹鍙栧緟妫€娴?docx锛屾彁鍙栦簡娈佃惤鏂囨湰銆佹钀芥牱寮忋€佸瓧浣撳瓧鍙枫€佸姞绮?鏂滀綋銆佺缉杩涖€佽璺濄€侀〉杈硅窛銆侀〉鐪夐〉鑴氳窛绂诲拰琛ㄦ牸棰勮绛夎瘉鎹€俓n"
                    "浣犱笉鑳借嚜宸辨墽琛?Python 浠ｇ爜锛涗綘蹇呴』鍩轰簬涓嬮潰鎻愪緵鐨?python-docx 瑙ｆ瀽璇佹嵁鍜岃鏂囧叏鏂囧唴瀹硅繘琛屾牸寮忓垽鏂€俓n"
                    "濡傛灉鏌愰」鏍煎紡鍦?python-docx 璇佹嵁涓己澶憋紝涓嶈兘鍑┖鍒ゆ柇涓洪敊璇紱鍙兘鏍囪涓轰綆缃俊搴﹀苟璁剧疆 need_manual_confirmation=true銆俓n\n"
                    "閿欒瀹氫綅寮哄埗瑕佹眰锛氭瘡鏉?issue 蹇呴』缁欏嚭鍙畾浣嶄綅缃€俓n"
                    "1. location 蹇呴』鏄璞★紝鑷冲皯鍖呭惈 page銆乻ection銆乸aragraph_index銆乷bject_type銆乹uote 浜斾釜瀛楁銆俓n"
                    "2. location.paragraph_index 蹇呴』浼樺厛浣跨敤 python-docx 瑙ｆ瀽璇佹嵁 paragraphs[].index 涓殑鐪熷疄娈佃惤搴忓彿锛涢〉杈硅窛绛夐〉闈㈢骇闂鏃犳硶瀵瑰簲娈佃惤鏃讹紝濉啓鏈€鎺ヨ繎鐨勭珷鑺傛爣棰樻钀藉簭鍙凤紝骞惰缃?object_type=\"page\"銆俓n"
                    "3. location.quote 蹇呴』鏄鏂囧師鏂囨垨 python-docx 璇佹嵁涓湡瀹炲瓨鍦ㄧ殑鐭枃鏈紝闀垮害 6 鍒?60 涓瓧绗︼紝涓嶈兘缂栭€犮€俓n"
                    "4. anchor_text 蹇呴』涓?location.quote 涓€鑷存垨鏄?location.quote 鐨勬洿鐭繛缁瓙涓诧紝鐢ㄤ簬鍦?Word 涓彃鍏ュ師鐢熸壒娉ㄣ€俓n"
                    "5. 涓嶅厑璁稿彧鍐欌€滃叏鏂団€濃€滈〉闈㈣缃€濃€滄湭鐭ヤ綅缃€濅綔涓哄敮涓€瀹氫綅锛涘鏋滅‘瀹炴棤娉曞畾浣嶏紝浠嶉渶濉啓鏈€鎺ヨ繎绔犺妭銆佹钀藉簭鍙凤紝骞跺皢 need_manual_confirmation=true銆俓n\n"
                    f"全文读取概况 JSON：\n{json.dumps(build_document_overview(checked_text, [checked_text]), ensure_ascii=False)}\n\n"
                    f"python-docx 结构化证据 JSON：\n{json.dumps(docx_evidence, ensure_ascii=False)}\n\n"
                    f"格式规则 JSON：\n{json.dumps(format_rule, ensure_ascii=False)}\n\n"
                    f"待检测论文文件名：{checked_file.name}\n\n论文全文内容：\n{checked_text}"
                ),
            },
        ],
        agent_name="agent2_check_paper_format",
        log_path=llm_log_output,
    )
    check_payload = parse_llm_json(check_content, "agent2_check_paper_format full paper")

    statistics = check_payload.get("statistics") or {}
    issues = [
        normalize_issue(issue, index)
        for index, issue in enumerate(check_payload.get("issues", []), start=1)
        if isinstance(issue, dict)
    ]
    statistics = {
        "total_issues": len(issues),
        "by_category": statistics.get("by_category") or {},
        "by_severity": statistics.get("by_severity") or {},
        "need_manual_confirmation": statistics.get("need_manual_confirmation", 0),
        "checked_rule_count": len(format_rule["rules"]),
    }
    check_result = {
        "paper_analysis": {
            "standard_file": format_rule_output.name,
            "paper_file": checked_file.name,
            "paper_type": "docx",
            "selected_parser": "python-docx + llm-document-reader",
            "llm_provider": settings["provider"],
            "llm_model": settings["model"],
            "check_summary": "已由 python-docx/XML 读取格式证据，并由 LLM 完成规则检查。",
            "checked_text_chars": len(checked_text),
            "checked_chunk_count": 1,
            "chunk_char_lengths": [len(checked_text)],
        },
        "statistics": statistics,
        "issues": issues,
        "llm_raw": {
            "chunk_count": 1,
            "full_paper": check_payload,
            "docx_evidence_summary": {
                "parser": docx_evidence.get("parser"),
                "paragraph_count": docx_evidence.get("paragraph_count"),
                "section_count": docx_evidence.get("section_count"),
                "table_count": docx_evidence.get("table_count"),
            },
        },
    }
    return format_rule, check_result, statistics


def run_hybrid_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    format_rule_output: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    format_rule, check_result, _ = run_llm_workflow(
        settings, standard_file, checked_file, format_rule_output, llm_log_output
    )
    return enrich_hybrid_result(format_rule, check_result, checked_file, "hybrid")


def run_compare_workflow(
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    format_rule_output: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    llm_format_rule, llm_check_result, llm_statistics = run_llm_workflow(
        settings, standard_file, checked_file, format_rule_output, llm_log_output
    )
    hybrid_format_rule = copy.deepcopy(llm_format_rule)
    hybrid_check_result = copy.deepcopy(llm_check_result)
    format_rule, hybrid_result, statistics = enrich_hybrid_result(
        hybrid_format_rule, hybrid_check_result, checked_file, "compare"
    )
    hybrid_result["compare_result"] = build_compare_result(llm_check_result, hybrid_result)
    hybrid_result["llm_raw"] = {
        **(hybrid_result.get("llm_raw") or {}),
        "workflow_comparison": {
            "llm_direct_statistics": llm_statistics,
            "hybrid_statistics": statistics,
        },
    }
    return format_rule, hybrid_result, statistics


def run_selected_workflow(
    workflow: str,
    settings: dict[str, Any],
    standard_file: Path,
    checked_file: Path,
    format_rule_output: Path,
    llm_log_output: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    if workflow == "base_format":
        return run_base_format_workflow(settings, standard_file, checked_file, llm_log_output)
    if workflow == "language_semantic":
        return run_language_semantic_workflow(settings, standard_file, checked_file, llm_log_output)
    if workflow == "full_check":
        return run_full_check_workflow(settings, standard_file, checked_file, llm_log_output)
    raise ValueError(f"Unsupported workflow: {workflow}")


def write_simple_docx(path: Path, lines: list[str]) -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ET.register_namespace("w", w_ns)

    def qn(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    document = ET.Element(qn("document"))
    body = ET.SubElement(document, qn("body"))
    for line in lines:
        paragraph = ET.SubElement(body, qn("p"))
        run = ET.SubElement(paragraph, qn("r"))
        text = ET.SubElement(run, qn("t"))
        text.text = line
    ET.SubElement(body, qn("sectPr"))

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    ensure_parent(path)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", content_types)
        package.writestr("_rels/.rels", rels)
        package.writestr("word/document.xml", ET.tostring(document, encoding="utf-8", xml_declaration=True))


def parse_required_xml_part(parts: dict[str, bytes], name: str) -> ET.Element:
    return ET.fromstring(parts[name])


def ensure_comments_part(parts: dict[str, bytes]) -> tuple[ET.Element, int]:
    if "word/comments.xml" in parts:
        root = ET.fromstring(parts["word/comments.xml"])
        ids = [
            int(comment.get(qn(W_NS, "id"), "0"))
            for comment in root.findall(qn(W_NS, "comment"))
            if str(comment.get(qn(W_NS, "id"), "0")).isdigit()
        ]
        return root, (max(ids) + 1 if ids else 0)
    return ET.Element(qn(W_NS, "comments")), 0


def ensure_document_comments_relationship(parts: dict[str, bytes]) -> None:
    rel_name = "word/_rels/document.xml.rels"
    root = ET.fromstring(parts[rel_name]) if rel_name in parts else ET.Element(qn(PKG_REL_NS, "Relationships"))
    for relationship in root.findall(qn(PKG_REL_NS, "Relationship")):
        if relationship.get("Type") == f"{OFFICE_REL_NS}/comments":
            parts[rel_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            return

    existing_ids = {relationship.get("Id", "") for relationship in root.findall(qn(PKG_REL_NS, "Relationship"))}
    index = 1
    while f"rId{index}" in existing_ids:
        index += 1
    ET.SubElement(
        root,
        qn(PKG_REL_NS, "Relationship"),
        {
            "Id": f"rId{index}",
            "Type": f"{OFFICE_REL_NS}/comments",
            "Target": "comments.xml",
        },
    )
    parts[rel_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)


def ensure_comments_content_type(parts: dict[str, bytes]) -> None:
    root = ET.fromstring(parts["[Content_Types].xml"])
    for override in root.findall(qn(CONTENT_TYPES_NS, "Override")):
        if override.get("PartName") == "/word/comments.xml":
            parts["[Content_Types].xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            return
    ET.SubElement(
        root,
        qn(CONTENT_TYPES_NS, "Override"),
        {
            "PartName": "/word/comments.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        },
    )
    parts["[Content_Types].xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)


def paragraph_text(paragraph: ET.Element) -> str:
    return "".join(node.text or "" for node in paragraph.iter(qn(W_NS, "t")))


def find_anchor_paragraph(document_root: ET.Element, issue: dict[str, Any]) -> ET.Element | None:
    anchor = str(issue.get("anchor_text") or issue.get("location", {}).get("quote") or "").strip()
    paragraphs = list(document_root.iter(qn(W_NS, "p")))
    if anchor:
        for paragraph in paragraphs:
            if anchor in paragraph_text(paragraph):
                return paragraph

    paragraph_index = issue.get("location", {}).get("paragraph_index")
    if paragraph_index not in (None, ""):
        try:
            index = int(paragraph_index) - 1
        except (TypeError, ValueError):
            return None
        if 0 <= index < len(paragraphs):
            return paragraphs[index]
    return None


def append_comment(comments_root: ET.Element, comment_id: int, text: str) -> None:
    comment = ET.SubElement(
        comments_root,
        qn(W_NS, "comment"),
        {
            qn(W_NS, "id"): str(comment_id),
            qn(W_NS, "author"): "LLM Format Check",
            qn(W_NS, "date"): datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        },
    )
    paragraph = ET.SubElement(comment, qn(W_NS, "p"))
    run = ET.SubElement(paragraph, qn(W_NS, "r"))
    text_node = ET.SubElement(run, qn(W_NS, "t"))
    text_node.text = text


def mark_paragraph_with_comment(paragraph: ET.Element, comment_id: int) -> None:
    start = ET.Element(qn(W_NS, "commentRangeStart"), {qn(W_NS, "id"): str(comment_id)})
    end = ET.Element(qn(W_NS, "commentRangeEnd"), {qn(W_NS, "id"): str(comment_id)})
    ref_run = ET.Element(qn(W_NS, "r"))
    ET.SubElement(ref_run, qn(W_NS, "commentReference"), {qn(W_NS, "id"): str(comment_id)})
    insert_at = 1 if len(paragraph) and paragraph[0].tag == qn(W_NS, "pPr") else 0
    paragraph.insert(insert_at, start)
    paragraph.append(end)
    paragraph.append(ref_run)


def write_annotated_copy(source: Path, target: Path, issues: list[dict[str, Any]] | None = None) -> None:
    ensure_parent(target)
    if not issues:
        if source.exists():
            shutil.copy2(source, target)
        else:
            write_simple_docx(target, ["未生成批注：源文件不存在。"])
        return
    if source.suffix.lower() != ".docx" or not source.exists():
        write_simple_docx(target, ["未生成批注：源文件不是 docx。"])
        return

    with zipfile.ZipFile(source, "r") as package:
        parts = {info.filename: package.read(info.filename) for info in package.infolist()}
    if "word/document.xml" not in parts:
        shutil.copy2(source, target)
        return

    document_root = parse_required_xml_part(parts, "word/document.xml")
    comments_root, next_id = ensure_comments_part(parts)
    placed = 0
    for issue in issues:
        paragraph = find_anchor_paragraph(document_root, issue)
        if paragraph is None:
            continue
        comment_id = next_id
        next_id += 1
        append_comment(comments_root, comment_id, str(issue.get("comment_text") or build_comment_text(issue)))
        mark_paragraph_with_comment(paragraph, comment_id)
        placed += 1

    if not placed:
        shutil.copy2(source, target)
        return

    parts["word/document.xml"] = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)
    parts["word/comments.xml"] = ET.tostring(comments_root, encoding="utf-8", xml_declaration=True)
    ensure_document_comments_relationship(parts)
    ensure_comments_content_type(parts)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        temp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as package:
            for name, content in parts.items():
                package.writestr(name, content)
        shutil.move(str(temp_path), target)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def build_file_link(label: str, file_type: str, path: Path) -> dict[str, str]:
    return {
        "label": label,
        "file_type": file_type,
        "path": str(path),
        "url": "/api/files/" + str(path).replace("\\", "/"),
    }


def run_request(request: dict[str, Any]) -> dict[str, Any]:
    work_dir = Path(request["work_dir"])
    checked_source = Path(request["checked_file"])
    standard_source = Path(request["standard_file"])
    workflow = normalize_workflow(request.get("workflow"))
    config_path = request.get("llm_config_path")
    local_config_path = request.get("llm_local_config_path")
    store = LlmConfigStore(Path(config_path), Path(local_config_path)) if config_path and local_config_path else None
    llm_settings = load_provider_settings(str(request.get("provider") or ""), str(request.get("model") or ""), store)

    default_paths = build_request_paths(work_dir, checked_source.name, standard_source.name)
    checked_file = copy_if_needed(checked_source, default_paths["checked_file"])
    standard_file = copy_if_needed(standard_source, default_paths["standard_file"])
    format_rule_output = Path(request.get("format_rule_output") or default_paths["format_rule_output"])
    check_result_output = Path(request.get("check_result_output") or default_paths["check_result_output"])
    analysis_output = Path(request.get("analysis_output") or default_paths["analysis_output"])
    analysis_markdown_output = default_paths["analysis_markdown_output"]
    llm_log_output = Path(request.get("llm_log_output") or default_paths["llm_log_output"])
    annotated_output = Path(request.get("annotated_output") or default_paths["annotated_output"])

    if llm_settings["api_key_configured"]:
        format_rule, check_result, statistics = run_selected_workflow(
            workflow, llm_settings, standard_file, checked_file, format_rule_output, llm_log_output
        )
    else:
        format_rule, check_result, statistics = build_placeholder_outputs(
            standard_file, checked_file, format_rule_output, workflow
        )

    write_json(format_rule_output, format_rule)
    write_json(check_result_output, check_result)
    write_simple_docx(
        analysis_output,
        [
            "论文格式检查错误统计分析",
            f"错误总数：{statistics.get('total_issues', 0)}",
            f"已检查规则数：{statistics.get('checked_rule_count', 0)}",
            f"需要人工确认：{statistics.get('need_manual_confirmation', 0)}",
        ],
    )
    ensure_parent(analysis_markdown_output)
    analysis_markdown_output.write_text(
        f"# 论文格式检查错误统计分析\n\n错误总数：{statistics.get('total_issues', 0)}\n",
        encoding="utf-8",
    )
    write_annotated_copy(checked_file, annotated_output, check_result.get("issues", []))

    return {
        "success": True,
        "message": "completed",
        "workflow": workflow,
        "format_rule": build_file_link("涓嬭浇鏍煎紡瑙勮寖 JSON", "format_rule", format_rule_output),
        "analysis_doc": build_file_link("涓嬭浇閿欒缁熻鍒嗘瀽鏂囨。", "analysis_doc", analysis_output),
        "annotated_doc": build_file_link("批注版论文", "annotated_paper", annotated_output),
        "statistics": statistics,
        "compare_result": check_result.get("compare_result"),
        "issues": check_result.get("issues", []),
    }


class AgentHttpHandler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def do_GET(self) -> None:
        if self.path == "/health":
            self.write_json(200, {"status": "ok", "service": "format-check-agent"})
            return
        self.write_json(404, {"success": False, "message": "not found"})

    def do_POST(self) -> None:
        if self.path != "/agent/check":
            self.write_json(404, {"success": False, "message": "not found"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            body = self.rfile.read(length).decode("utf-8")
            request_payload = json.loads(body)
            result = run_request(request_payload)
            self.write_json(200, result)
        except Exception as exc:
            self.write_json(500, {"success": False, "message": str(exc), "data": None})

    def write_json(self, status: int, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: Any) -> None:
        print(f"[agent-service] {self.address_string()} - {format % args}", file=sys.stderr, flush=True)


def run_server(host: str, port: int) -> None:
    server = ThreadingHTTPServer((host, port), AgentHttpHandler)
    print(f"[agent-service] Python Agent \u5df2\u542f\u52a8\uff1ahttp://{host}:{port}", file=sys.stderr, flush=True)
    print("[agent-service] Java \u540e\u7aef\u5c06\u901a\u8fc7 POST /agent/check \u8c03\u7528\u5f53\u524d\u5e38\u9a7b\u670d\u52a1", file=sys.stderr, flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("[agent-service] \u6536\u5230\u505c\u6b62\u4fe1\u53f7\uff0c\u6b63\u5728\u5173\u95ed", file=sys.stderr, flush=True)
    finally:
        server.server_close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run local format-check agent workflow.")
    parser.add_argument("--request", help="Path to request JSON generated by Java.")
    parser.add_argument("--serve", action="store_true", help="Run as a long-lived local HTTP agent service.")
    parser.add_argument("--host", default="127.0.0.1", help="Agent service host.")
    parser.add_argument("--port", type=int, default=8001, help="Agent service port.")
    args = parser.parse_args(argv)
    if args.request:
        with open(args.request, encoding="utf-8") as request_file:
            request = json.load(request_file)
        result = run_request(request)
        print(json.dumps(result, ensure_ascii=False))
        return 0

    run_server(args.host, args.port)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
