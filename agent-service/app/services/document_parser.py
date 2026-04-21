from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.models.contracts import ParsedDocument


class DocumentParser:
    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        suffix = path.suffix.lower()
        if suffix == ".docx":
            return self._parse_docx(path)
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix in {".md", ".txt", ".js", ".json"}:
            return self._parse_text(path)
        raise ValueError(f"Unsupported file type: {suffix}")

    def _parse_docx(self, path: Path) -> ParsedDocument:
        document = Document(path)
        structure_lines: list[str] = []
        plain_parts: list[str] = []
        excerpts: list[str] = []

        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = (paragraph.text or "").strip()
            if not text:
                continue

            run = paragraph.runs[0] if paragraph.runs else None
            font_name = run.font.name if run and run.font and run.font.name else ""
            font_size = round(run.font.size.pt, 1) if run and run.font and run.font.size else None
            alignment = str(paragraph.alignment).split(".")[-1] if paragraph.alignment is not None else "UNKNOWN"
            style_name = paragraph.style.name if paragraph.style else ""

            structure_lines.append(
                f"P{index} | style={style_name or 'Unknown'} | align={alignment} | font={font_name or 'Unknown'} | size={font_size or 'Unknown'} | text={text}"
            )
            plain_parts.append(text)
            if len(excerpts) < 30:
                excerpts.append(text)

        metadata = {
            "paragraphCount": len(document.paragraphs),
            "sectionCount": len(document.sections),
            "coreProperties": {
                "title": document.core_properties.title,
                "author": document.core_properties.author,
            },
        }
        return ParsedDocument.from_parts(
            file_path=path,
            file_type="docx",
            plain_text="\n".join(plain_parts),
            structure_lines=structure_lines,
            excerpts=excerpts,
            metadata=metadata,
        )

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        reader = PdfReader(str(path))
        structure_lines: list[str] = []
        plain_parts: list[str] = []
        excerpts: list[str] = []

        for page_index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            compact = " ".join(text.split())
            structure_lines.append(f"Page {page_index}: {compact[:1000]}")
            plain_parts.append(text)
            if len(excerpts) < 20:
                excerpts.append(compact[:800])

        metadata = {
            "pageCount": len(reader.pages),
            "pdfMetadata": dict(reader.metadata or {}),
        }
        return ParsedDocument.from_parts(
            file_path=path,
            file_type="pdf",
            plain_text="\n\n".join(plain_parts),
            structure_lines=structure_lines,
            excerpts=excerpts,
            metadata=metadata,
        )

    def _parse_text(self, path: Path) -> ParsedDocument:
        raw = path.read_text(encoding="utf-8")
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        structure_lines = [f"L{index}: {line}" for index, line in enumerate(lines[:300], start=1)]
        excerpts = lines[:30]
        return ParsedDocument.from_parts(
            file_path=path,
            file_type=path.suffix.lower().lstrip("."),
            plain_text=raw,
            structure_lines=structure_lines,
            excerpts=excerpts,
            metadata={"lineCount": len(lines)},
        )
