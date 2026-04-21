from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

from app.models.contracts import CheckIssue


class AnnotationWriter:
    def write(self, source_file_path: str, annotated_file_path: str, issues: list[CheckIssue]) -> str:
        source = Path(source_file_path)
        target = Path(annotated_file_path)
        target.parent.mkdir(parents=True, exist_ok=True)

        if source.suffix.lower() == ".docx":
            self._write_docx_annotations(source, target, issues)
        else:
            shutil.copyfile(source, target)
        return str(target)

    def _write_docx_annotations(self, source: Path, target: Path, issues: list[CheckIssue]) -> None:
        document = Document(source)
        document.add_page_break()
        document.add_heading("格式检查批注", level=1)

        if not issues:
            document.add_paragraph("未发现格式问题。")
        else:
            for issue in issues:
                document.add_paragraph(
                    f"{issue.id} | 位置: 第{issue.location.page or '-'}页 / 段落{issue.location.paragraphIndex or '-'} / "
                    f"{issue.location.heading or '-'} | 规则: {issue.ruleName} | 原因: {issue.reason} | 建议: {issue.suggestion}"
                )

        document.save(target)
